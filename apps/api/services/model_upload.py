"""Replication Studio (rep3) — reading an uploaded model file. Pure parsing and
validation; nothing here writes to the database or to R2 (services.models does).

UPLOADED FILES ARE UNTRUSTED
──────────────────────────────────────────────────────────────────────────────
* Accepted: ``.xlsx`` (openpyxl, read-only, ``data_only=True``), ``.xls``
  (xlrd — reads the BIFF records, cached values only) and ``.csv`` (UTF-8 or
  UTF-8 with BOM, delimiter sniffed among , ; tab |). Refused: ``.xlsm`` /
  ``.xlsb`` / ``.xltm`` (and any .xlsx whose package carries a VBA project or a
  macro-enabled content type), anything else, anything over 5 MB, and any
  sheet with more than 5,000 data rows.
* Formulas are NEVER evaluated. openpyxl only ever reports the value Excel
  cached when it last saved the file; a formula cell with no cached value is a
  row-level validation error. (The workbook is opened twice, both read-only:
  once for cached values, once to learn which cells hold a formula.) Macros are
  never run — nothing here can run one.
* An .xlsx is a zip: before openpyxl sees it, every member is checked against
  a total-expanded cap (50 MB) and a per-member expansion ratio (100×) — a
  decompression bomb is refused, not opened. XML entity expansion is bounded by
  the interpreter's expat (2.4.1+ refuses billion-laughs amplification).

COLUMNS. A server vocabulary of synonyms (:data:`FIELDS`) suggests the mapping
from each header; a provider's saved mapping (services.models) is applied
first; the user confirms. The header row is the first row (of the first 30)
whose cells name an identifier and a weight — or the row the user chooses.

WEIGHTS are Decimal end to end, never float: a spreadsheet number is read
through ``str()`` of the cell value (the shortest repr), a text cell has its %
and thousands separators removed. Percent (sums ~100) or fraction (sums ~1) is
detected, stated, and converted to a fraction.

IDENTIFIERS resolve through the ONE security master
(``services.securities_global.get_by_identifiers_bulk`` — merge-following,
ambiguous values absent). A row that gives a CUSIP or ISIN is resolved by those
ONLY: if they do not resolve the row is unresolved — it is never matched on its
ticker. A ticker (or SEDOL) is used only when the row gives neither, and only
when it names exactly one security. Unresolved rows are queued on the version
(resolution_state 'unresolved', raw identifiers kept) — never guessed.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

MAX_BYTES = 5 * 1024 * 1024
MAX_DATA_ROWS = 5_000
MAX_ZIP_TOTAL = 50 * 1024 * 1024
MAX_ZIP_RATIO = 100
HEADER_SCAN_ROWS = 30
DEFAULT_TOLERANCE_PCT = Decimal("0.05")          # weights must sum to 100% ± this many percent
TOLERANCE_SETTING_KEY = "portfolio.models.weight_tolerance_pct"
DIFF_THRESHOLD = Decimal("0.0001")               # a weight move of more than 1 bp is listed in the diff
Q = Decimal("0.0000000001")                      # numeric(14, 10)

ACCEPTED_EXTENSIONS = (".xlsx", ".xls", ".csv")
MACRO_EXTENSIONS = (".xlsm", ".xlsb", ".xltm", ".xlam")
CSV_SHEET = "CSV"

# key → (label, synonyms). The label is the template's header; synonyms are
# compared after normalisation (lower case, every non-alphanumeric run → one space).
FIELDS: dict[str, tuple[str, tuple[str, ...]]] = {
    "ticker": ("Ticker", ("ticker", "symbol", "ticker symbol", "trading symbol")),
    "cusip": ("CUSIP", ("cusip", "cusip number")),
    "isin": ("ISIN", ("isin", "isin code")),
    "sedol": ("SEDOL", ("sedol",)),
    "name": ("Name", ("name", "security name", "security", "description", "holding", "company", "security description")),
    "weight": ("Target Weight", ("weight", "target", "target weight", "target", "allocation", "target allocation",
                                 "model weight", "weight pct", "pct", "percent", "portfolio weight", "target pct")),
    "asset_class": ("Asset Class", ("asset class", "class", "asset type", "asset category")),
    "sleeve": ("Sleeve", ("sleeve", "sleeve name", "strategy")),
    "drift_min": ("Min", ("min", "minimum", "min weight", "lower band", "band min", "min pct")),
    "drift_max": ("Max", ("max", "maximum", "max weight", "upper band", "band max", "max pct")),
    "alternates": ("Alternate", ("alternate", "alternates", "substitute", "alternate ticker", "alternative")),
    "restrictions": ("Restriction", ("restriction", "restrictions", "restricted")),
    "notes": ("Notes", ("notes", "note", "comment", "comments")),
    "cash": ("Cash", ("cash", "is cash", "cash flag")),
}
IDENTIFIER_FIELDS = ("ticker", "cusip", "isin", "sedol")
CASH_WORDS = {"cash", "$cash", "cash usd", "usd cash", "cash and equivalents", "cash equivalents", "cash usd"}
TRUTHY = {"y", "yes", "true", "1", "x", "cash"}

SEVERITY_LABELS = {"error": "Error — blocks the draft", "warning": "Warning — blocks approval",
                   "info": "Note"}
VALIDATION_CODES = {
    "file_type": "File type not accepted",
    "macro_refused": "Macro-enabled workbook refused",
    "too_large": "File larger than 5 MB",
    "too_many_rows": "More than 5,000 rows",
    "archive_ratio": "Workbook expands beyond a safe size",
    "unreadable": "File could not be read",
    "reader_unavailable": "Reader not installed on this server",
    "encoding": "Text encoding not UTF-8",
    "no_header": "No header row found",
    "mapping": "Column mapping incomplete",
    "formula_no_value": "Formula with no saved value",
    "error_cell": "Spreadsheet error value",
    "no_identifier": "Row has no identifier",
    "weight_missing": "Weight missing",
    "weight_not_number": "Weight is not a number",
    "weight_negative": "Negative weight",
    "weight_range": "Weight above 100%",
    "weight_sum": "Weights do not sum to 100%",
    "weight_basis": "Weight basis",
    "cash_multiple": "More than one cash row",
    "cash_filled": "Gap filled with cash",
    "cash_remainder": "Cash target set to the remainder",
    "duplicate_security": "Security listed twice",
    "duplicate_identifier": "Identifier listed twice",
    "unresolved": "Identifier not resolved (queued)",
    "identifier_conflict": "Identifiers point at different securities",
    "band_range": "Drift band outside 0–100%",
    "band_order": "Drift band does not contain the target",
    "band_not_number": "Drift band is not a number",
    "alternate_unresolved": "Alternate not resolved",
    "total_row": "Total row ignored",
    "tradability": "Tradability note",
    "no_rows": "No positions",
}


class UploadRefused(Exception):
    """The file itself is refused (type, macro, size, rows, archive, unreadable)."""

    def __init__(self, code: str, text: str):
        super().__init__(text)
        self.code = code
        self.text = text


@dataclass
class Message:
    code: str
    severity: str
    text: str
    row: int | None = None
    column: str | None = None

    def as_dict(self) -> dict:
        return {"code": self.code, "severity": self.severity, "row": self.row, "column": self.column,
                "text": self.text}


@dataclass
class Workbook:
    kind: str
    sheets: list[str]
    grids: dict[str, list[list[Any]]]
    #: (sheet, row_num, col_index) of formula cells with no cached value
    formula_gaps: set = field(default_factory=set)
    #: (sheet, row_num, col_index) of spreadsheet error values (#N/A, #REF! …)
    error_cells: set = field(default_factory=set)


# ── normalisation ───────────────────────────────────────────────────────────
def norm_header(v) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(v or "").lower()).strip()


_SYNONYMS: dict[str, str] = {}
for _k, (_label, _syns) in FIELDS.items():
    for _s in _syns:
        _SYNONYMS.setdefault(norm_header(_s), _k)
    _SYNONYMS.setdefault(norm_header(_label), _k)
for _alias in ("target %", "weight %", "% weight", "target weight %", "allocation %", "% of portfolio"):
    _SYNONYMS.setdefault(norm_header(_alias), "weight")


_VOCABS: dict[int, dict[str, str]] = {}


def synonyms_for(fields: dict | None) -> dict[str, str]:
    """The normalised-header → field map of a field vocabulary (rep7: the
    held-away positions-and-lots file reuses this machinery with its own
    vocabulary; ``None`` = the model vocabulary, unchanged)."""
    if fields is None or fields is FIELDS:
        return _SYNONYMS
    got = _VOCABS.get(id(fields))
    if got is None:
        got = {}
        for k, (label, syns) in fields.items():
            for s_ in syns:
                got.setdefault(norm_header(s_), k)
            got.setdefault(norm_header(label), k)
        _VOCABS[id(fields)] = got
    return got


def field_for_header(header, fields: dict | None = None) -> str | None:
    return synonyms_for(fields).get(norm_header(header))


def cell_text(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return str(int(v))
        return repr(v)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    s = str(v).strip()
    return s or None


def to_decimal(v) -> Decimal | None:
    """A cell as an exact Decimal, or None when blank. Raises ValueError when
    it is not a number. Never goes through float arithmetic: a float cell is
    read via its shortest repr."""
    if v is None:
        return None
    if isinstance(v, bool):
        raise ValueError("a yes/no value, not a number")
    if isinstance(v, Decimal):
        return v
    if isinstance(v, int):
        return Decimal(v)
    if isinstance(v, float):
        return Decimal(repr(v))
    s = str(v).strip().replace(" ", "").replace(" ", "")
    if not s:
        return None
    s = s.rstrip("%")
    if re.fullmatch(r"[-+]?\d{1,3}(,\d{3})+(\.\d+)?", s):
        s = s.replace(",", "")
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError(f"{str(v)!r} is not a number") from None
    if not d.is_finite():
        raise ValueError(f"{str(v)!r} is not a number")
    return d


def has_percent_sign(v) -> bool:
    return isinstance(v, str) and v.strip().endswith("%")


def q(d: Decimal) -> Decimal:
    return d.quantize(Q)


def pct_text(d: Decimal) -> str:
    return f"{(d * 100).quantize(Decimal('0.0001'))}%"


def csv_safe(v) -> str:
    """A value for a CSV export: a leading = + - @ tab or CR (a formula prefix
    a spreadsheet would act on) is neutralised with a leading apostrophe."""
    if v is None:
        return ""
    s = str(v)
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + s
    return s


# ── reading the file ────────────────────────────────────────────────────────
def extension(filename: str) -> str:
    name = (filename or "").strip().lower()
    return name[name.rfind("."):] if "." in name else ""


def precheck(data: bytes, filename: str) -> str:
    """The file kind, or UploadRefused. Runs before anything opens the file."""
    ext = extension(filename)
    if ext in MACRO_EXTENSIONS:
        raise UploadRefused("macro_refused", f"{ext} files can carry macros and are refused — save the workbook "
                                             "as .xlsx (or .csv) and upload that")
    if ext not in ACCEPTED_EXTENSIONS:
        raise UploadRefused("file_type", f"{ext or 'this file type'} is not accepted — upload .xlsx, .xls or .csv")
    if len(data) > MAX_BYTES:
        raise UploadRefused("too_large", f"the file is {len(data):,} bytes; the limit is {MAX_BYTES:,} (5 MB)")
    if not data:
        raise UploadRefused("unreadable", "the file is empty")
    if ext == ".xlsx":
        if not data.startswith(b"PK\x03\x04"):
            raise UploadRefused("unreadable", "this .xlsx is not an Excel workbook (not a zip package)")
        _check_zip(data)
        return "xlsx"
    if ext == ".xls":
        if not data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            raise UploadRefused("unreadable", "this .xls is not a legacy Excel workbook")
        return "xls"
    return "csv"


def _check_zip(data: bytes) -> None:
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        raise UploadRefused("unreadable", "the workbook package is damaged") from None
    with zf:
        total = 0
        for info in zf.infolist():
            name = info.filename.lower()
            if name.endswith("vbaproject.bin") or name.endswith(".bin") and "vba" in name:
                raise UploadRefused("macro_refused", "the workbook carries a VBA macro project and is refused")
            total += info.file_size
            if total > MAX_ZIP_TOTAL:
                raise UploadRefused("archive_ratio", "the workbook expands beyond 50 MB and is refused")
            if info.file_size > 1_000_000 and info.file_size > MAX_ZIP_RATIO * max(info.compress_size, 1):
                raise UploadRefused("archive_ratio", f"a part of the workbook ({info.filename}) expands more than "
                                                     f"{MAX_ZIP_RATIO}× and is refused")
        try:
            ctypes = zf.read("[Content_Types].xml") if "[Content_Types].xml" in zf.namelist() else b""
        except (KeyError, zipfile.BadZipFile, RuntimeError):
            ctypes = b""
        if b"macroEnabled" in ctypes or b"vbaProject" in ctypes:
            raise UploadRefused("macro_refused", "the workbook is macro-enabled and is refused")


def read_workbook(data: bytes, filename: str) -> Workbook:
    kind = precheck(data, filename)
    if kind == "xlsx":
        return _read_xlsx(data)
    if kind == "xls":
        return _read_xls(data)
    return _read_csv(data)


def _cap_rows(n: int, sheet: str) -> None:
    if n > MAX_DATA_ROWS + HEADER_SCAN_ROWS:
        raise UploadRefused("too_many_rows", f"sheet {sheet!r} has more than {MAX_DATA_ROWS:,} rows — split the file")


def _read_xlsx(data: bytes) -> Workbook:
    import openpyxl

    try:
        values_wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        formulas_wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except UploadRefused:
        raise
    except Exception as exc:  # noqa: BLE001 — any parser failure is "unreadable", never a 500
        raise UploadRefused("unreadable", f"the workbook could not be read ({type(exc).__name__})") from None
    try:
        sheets = list(values_wb.sheetnames)
        grids: dict[str, list[list[Any]]] = {}
        gaps: set = set()
        errors: set = set()
        for name in sheets:
            rows: list[list[Any]] = []
            for i, row in enumerate(values_wb[name].iter_rows(min_row=1, values_only=True), start=1):
                rows.append(list(row))
                _cap_rows(i, name)
            for i, row in enumerate(formulas_wb[name].iter_rows(min_row=1, values_only=True), start=1):
                if i > len(rows) + 1:
                    break
                for j, v in enumerate(row):
                    if isinstance(v, str) and v.startswith("="):
                        cached = rows[i - 1][j] if i - 1 < len(rows) and j < len(rows[i - 1]) else None
                        if cached is None:
                            gaps.add((name, i, j))
            for i, row in enumerate(rows, start=1):
                for j, v in enumerate(row):
                    if isinstance(v, str) and v in ("#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"):
                        errors.add((name, i, j))
            grids[name] = rows
    finally:
        values_wb.close()
        formulas_wb.close()
    return Workbook("xlsx", sheets, grids, gaps, errors)


def _read_xls(data: bytes) -> Workbook:
    try:
        import xlrd
    except ImportError:
        raise UploadRefused("reader_unavailable", "the .xls reader (xlrd) is not installed on this server — "
                                                  "save the file as .xlsx or .csv") from None
    try:
        book = xlrd.open_workbook(file_contents=data, on_demand=True, formatting_info=False)
    except Exception as exc:  # noqa: BLE001
        raise UploadRefused("unreadable", f"the workbook could not be read ({type(exc).__name__})") from None
    try:
        sheets = list(book.sheet_names())
        grids: dict[str, list[list[Any]]] = {}
        errors: set = set()
        for name in sheets:
            sh = book.sheet_by_name(name)
            _cap_rows(sh.nrows, name)
            rows = []
            for i in range(sh.nrows):
                out = []
                for j in range(sh.ncols):
                    c = sh.cell(i, j)
                    if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        out.append(None)
                    elif c.ctype == xlrd.XL_CELL_ERROR:
                        out.append(None)
                        errors.add((name, i + 1, j))
                    elif c.ctype == xlrd.XL_CELL_BOOLEAN:
                        out.append(bool(c.value))
                    else:
                        out.append(c.value)
                rows.append(out)
            grids[name] = rows
            book.unload_sheet(name)
    finally:
        book.release_resources()
    return Workbook("xls", sheets, grids, set(), errors)


def _read_csv(data: bytes) -> Workbook:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise UploadRefused("encoding", "the CSV is not UTF-8 — save it as 'CSV UTF-8' and upload again") from None
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = ","
    rows = []
    for i, rec in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        _cap_rows(i, CSV_SHEET)
        rows.append([c if c.strip() else None for c in rec])
    return Workbook("csv", [CSV_SHEET], {CSV_SHEET: rows})


# ── header + mapping ────────────────────────────────────────────────────────
def detect_header(grid: list[list[Any]], saved: dict | None = None, *, fields: dict | None = None,
                  required: tuple[str, ...] = ("weight",)) -> int | None:
    """1-based number of the first row (of the first 30) that names an
    identifier column and every ``required`` column (a weight, for a model) —
    by the saved provider mapping first, then the synonym vocabulary (a
    provider's own headers are often in neither list until its mapping is
    saved)."""
    vocab = fields or FIELDS
    saved_norm = {norm_header(k): v for k, v in (saved or {}).items() if v in vocab}
    for i, row in enumerate(grid[:HEADER_SCAN_ROWS], start=1):
        found = {saved_norm.get(norm_header(c)) or field_for_header(c, fields) for c in row if cell_text(c)}
        if set(required) <= found and found & set(IDENTIFIER_FIELDS):
            return i
    return None


def headers_of(grid: list[list[Any]], header_row: int) -> list[str]:
    if header_row < 1 or header_row > len(grid):
        return []
    return [cell_text(c) or "" for c in grid[header_row - 1]]


def suggest_mapping(headers: list[str], saved: dict | None = None, *, fields: dict | None = None) -> list[dict]:
    """One suggestion per column: the saved provider mapping first (keyed by
    normalised header), then the synonym vocabulary. Each field once."""
    saved_norm = {norm_header(k): v for k, v in (saved or {}).items() if v in (fields or FIELDS)}
    used: set[str] = set()
    out = []
    for idx, h in enumerate(headers):
        key = norm_header(h)
        fld, src = None, None
        if key and key in saved_norm and saved_norm[key] not in used:
            fld, src = saved_norm[key], "saved"
        elif key:
            cand = field_for_header(h, fields)
            if cand and cand not in used:
                fld, src = cand, "synonym"
        if fld:
            used.add(fld)
        out.append({"column": idx, "header": h, "field": fld, "source": src})
    return out


def mapping_from_request(headers: list[str], mapping: dict | None, saved: dict | None = None, *,
                         fields: dict | None = None) -> dict[int, str]:
    """{column index: field}. ``mapping`` is {header text: field | ''}; a header
    it does not mention takes the suggestion."""
    vocab = fields or FIELDS
    sugg = {s["column"]: s["field"] for s in suggest_mapping(headers, saved, fields=fields) if s["field"]}
    if mapping is None:
        return sugg
    by_norm = {norm_header(k): (v or None) for k, v in mapping.items()}
    out: dict[int, str] = {}
    for idx, h in enumerate(headers):
        key = norm_header(h)
        fld = by_norm[key] if key in by_norm else sugg.get(idx)
        if fld:
            if fld not in vocab:
                raise UploadRefused("mapping", f"unknown field {fld!r} for column {h!r}")
            out[idx] = fld
    seen: dict[str, int] = {}
    for idx, fld in out.items():
        if fld in seen:
            raise UploadRefused("mapping", f"two columns are mapped to {vocab[fld][0]!r} "
                                           f"({headers[seen[fld]]!r} and {headers[idx]!r})")
        seen[fld] = idx
    return out


def mapping_by_header(headers: list[str], cols: dict[int, str]) -> dict[str, str]:
    return {headers[i]: f for i, f in sorted(cols.items()) if i < len(headers) and headers[i]}


# ── interpreting rows ───────────────────────────────────────────────────────
@dataclass
class Row:
    row_num: int
    raw: dict[str, Any]
    ticker: str | None = None
    cusip: str | None = None
    isin: str | None = None
    sedol: str | None = None
    name: str | None = None
    weight: Decimal | None = None
    asset_class: str | None = None
    sleeve: str | None = None
    drift_min: Decimal | None = None
    drift_max: Decimal | None = None
    alternates: list = field(default_factory=list)
    restrictions: str | None = None
    notes: str | None = None
    is_cash: bool = False
    global_security_id: str | None = None
    security_name: str | None = None
    state: str = "unresolved"
    resolution_note: str | None = None

    def key(self) -> str:
        if self.is_cash:
            return "cash"
        if self.global_security_id:
            return str(self.global_security_id)
        ident = self.cusip or self.isin or self.sedol or self.ticker or self.name or f"row{self.row_num}"
        return f"unresolved:{ident.upper()}"

    def label(self) -> str:
        if self.is_cash:
            return "Cash"
        return self.ticker or self.security_name or self.name or self.cusip or self.isin or self.sedol or "?"

    def as_position(self) -> dict:
        return {
            "row_num": self.row_num, "global_security_id": self.global_security_id, "raw_ticker": self.ticker,
            "raw_cusip": self.cusip, "raw_isin": self.isin, "raw_sedol": self.sedol,
            "raw_name": self.name or self.security_name, "target_weight": q(self.weight or Decimal(0)),
            "asset_class": self.asset_class, "sleeve_label": self.sleeve,
            "drift_min": q(self.drift_min) if self.drift_min is not None else None,
            "drift_max": q(self.drift_max) if self.drift_max is not None else None,
            "alternates": self.alternates, "restrictions": self.restrictions, "notes": self.notes,
            "resolution_state": "cash" if self.is_cash else self.state, "resolution_note": self.resolution_note,
        }


@dataclass
class ParseResult:
    sheet: str
    header_row: int
    headers: list[str]
    mapping: dict[str, str]
    weight_basis: str | None
    rows: list[Row]
    messages: list[Message]
    cash_target: Decimal | None = None
    tolerance: Decimal = DEFAULT_TOLERANCE_PCT

    @property
    def blocking(self) -> bool:
        return any(m.severity == "error" for m in self.messages)

    @property
    def counts(self) -> dict:
        return {"read": len(self.rows), "resolved": sum(1 for r in self.rows if r.state == "resolved" and not r.is_cash),
                "unresolved": sum(1 for r in self.rows if r.state == "unresolved" and not r.is_cash),
                "cash": sum(1 for r in self.rows if r.is_cash)}

    def positions(self) -> list[dict]:
        return [r.as_position() for r in self.rows]

    def report(self) -> dict:
        return {"sheet": self.sheet, "header_row": self.header_row, "headers": self.headers, "mapping": self.mapping,
                "weight_basis": self.weight_basis, "counts": self.counts,
                "cash_target": str(self.cash_target) if self.cash_target is not None else None,
                "tolerance_pct": str(self.tolerance), "blocking": self.blocking,
                "messages": [m.as_dict() for m in self.messages],
                "positions": [{**p, "target_weight": str(p["target_weight"]),
                               "drift_min": str(p["drift_min"]) if p["drift_min"] is not None else None,
                               "drift_max": str(p["drift_max"]) if p["drift_max"] is not None else None,
                               "global_security_id": str(p["global_security_id"]) if p["global_security_id"] else None}
                              for p in self.positions()]}


def _is_truthy(v) -> bool:
    return (cell_text(v) or "").lower() in TRUTHY


def extract_rows(wb: Workbook, sheet: str, header_row: int, cols: dict[int, str]) -> tuple[list[Row], list[Message]]:
    grid = wb.grids[sheet]
    msgs: list[Message] = []
    rows: list[Row] = []
    by_field = {f: i for i, f in cols.items()}
    for i in range(header_row, len(grid)):
        row_num = i + 1
        cells = grid[i]
        raw = {f: (cells[c] if c < len(cells) else None) for f, c in by_field.items()}
        if all(cell_text(v) is None for v in raw.values()):
            gap_cols = [c for c in by_field.values() if (sheet, row_num, c) in wb.formula_gaps]
            if not gap_cols:
                continue
        label_text = " ".join((cell_text(raw.get(f)) or "").lower() for f in ("ticker", "name", "cusip"))
        if re.search(r"\b(grand )?total\b", label_text) and not any(cell_text(raw.get(f)) for f in ("isin", "sedol")):
            msgs.append(Message("total_row", "info", f"row {row_num}: a total row is ignored", row_num))
            continue
        if len(rows) >= MAX_DATA_ROWS:
            raise UploadRefused("too_many_rows", f"more than {MAX_DATA_ROWS:,} data rows — split the file")
        r = Row(row_num=row_num, raw={k: cell_text(v) for k, v in raw.items()})
        for f, c in by_field.items():
            if (sheet, row_num, c) in wb.formula_gaps:
                msgs.append(Message("formula_no_value", "error",
                                    f"row {row_num}: the {FIELDS[f][0]} cell is a formula with no saved value — "
                                    "formulas are never evaluated; open the file in Excel, save it, and upload again",
                                    row_num, f))
            if (sheet, row_num, c) in wb.error_cells:
                msgs.append(Message("error_cell", "error", f"row {row_num}: the {FIELDS[f][0]} cell holds a "
                                                           "spreadsheet error value", row_num, f))
        r.ticker = (cell_text(raw.get("ticker")) or "").upper() or None
        r.cusip = (cell_text(raw.get("cusip")) or "").upper() or None
        if r.cusip and isinstance(raw.get("cusip"), (int, float)) and r.cusip.isdigit() and len(r.cusip) < 9:
            r.cusip = r.cusip.zfill(9)      # a numeric CUSIP cell lost its leading zeros
        r.isin = (cell_text(raw.get("isin")) or "").upper() or None
        r.sedol = (cell_text(raw.get("sedol")) or "").upper() or None
        r.name = cell_text(raw.get("name"))
        r.asset_class = cell_text(raw.get("asset_class"))
        r.sleeve = cell_text(raw.get("sleeve"))
        r.restrictions = cell_text(raw.get("restrictions"))
        r.notes = cell_text(raw.get("notes"))
        alt = cell_text(raw.get("alternates"))
        r.alternates = [a.strip().upper() for a in re.split(r"[;,|/]", alt) if a.strip()] if alt else []
        words = {(r.ticker or "").lower(), (r.name or "").lower(), (r.cusip or "").lower()} - {""}
        r.is_cash = _is_truthy(raw.get("cash")) or bool(words & CASH_WORDS)
        rows.append(r)
    return rows, msgs


def interpret_weights(rows: list[Row], *, header_row: int, tolerance_pct: Decimal,
                      fill_with_cash: bool, basis: str | None = None) -> tuple[str | None, Decimal | None, list[Message]]:
    """Detect percent vs fraction, convert to fractions, check the sum.
    Returns (basis, cash_target, messages). Mutates the rows' weights/bands."""
    msgs: list[Message] = []
    raw_w: dict[int, Decimal | None] = {}
    percent_sign = False
    for r in rows:
        v = r.raw.get("weight")
        percent_sign = percent_sign or has_percent_sign(v)
        try:
            raw_w[r.row_num] = to_decimal(v)
        except ValueError as exc:
            msgs.append(Message("weight_not_number", "error", f"row {r.row_num}: weight {exc}", r.row_num, "weight"))
            raw_w[r.row_num] = None
            continue
        if raw_w[r.row_num] is None and not r.is_cash:
            msgs.append(Message("weight_missing", "error", f"row {r.row_num}: no weight", r.row_num, "weight"))
    total_raw = sum((w for w in raw_w.values() if w is not None), Decimal(0))
    if basis is None:
        if percent_sign:
            basis = "percent"
        elif not rows:
            basis = None
        else:
            basis = "percent" if abs(total_raw - 100) < abs(total_raw - 1) else "fraction"
    scale = Decimal(100) if basis == "percent" else Decimal(1)
    if basis:
        msgs.append(Message("weight_basis", "info",
                            f"weights read as {'percentages (they sum to about 100)' if basis == 'percent' else 'fractions (they sum to about 1)'}"
                            " and stored as fractions", header_row, "weight"))
    for r in rows:
        w = raw_w.get(r.row_num)
        r.weight = (w / scale) if w is not None else None
        if r.weight is not None:
            if r.weight < 0:
                msgs.append(Message("weight_negative", "error", f"row {r.row_num}: negative weight "
                                                                f"{pct_text(r.weight)} — long/short models are not "
                                                                "supported yet", r.row_num, "weight"))
            elif r.weight > 1:
                msgs.append(Message("weight_range", "error", f"row {r.row_num}: weight {pct_text(r.weight)} is above "
                                                             "100%", r.row_num, "weight"))
        for f in ("drift_min", "drift_max"):
            try:
                b = to_decimal(r.raw.get(f))
            except ValueError as exc:
                msgs.append(Message("band_not_number", "error", f"row {r.row_num}: {FIELDS[f][0]} {exc}", r.row_num, f))
                continue
            setattr(r, f, (b / scale) if b is not None else None)
    cash_rows = [r for r in rows if r.is_cash]
    if len(cash_rows) > 1:
        msgs.append(Message("cash_multiple", "error", f"row {cash_rows[1].row_num}: a second cash row (the first is "
                                                      f"row {cash_rows[0].row_num})", cash_rows[1].row_num))
    tol = tolerance_pct / 100
    invested = sum((r.weight for r in rows if r.weight is not None and not r.is_cash), Decimal(0))
    cash_row = cash_rows[0] if cash_rows else None
    cash_target = None
    if cash_row is not None and cash_row.weight is None:
        gap = 1 - invested
        if gap < -tol:
            msgs.append(Message("weight_sum", "error", f"weights sum to {pct_text(invested)} before cash — more than "
                                                       f"100% (±{tolerance_pct}%)", header_row, "weight"))
        else:
            cash_row.weight = max(gap, Decimal(0))
            msgs.append(Message("cash_remainder", "info", f"row {cash_row.row_num}: the cash target is the remainder, "
                                                          f"{pct_text(cash_row.weight)}", cash_row.row_num, "weight"))
    total = invested + (cash_row.weight if cash_row is not None and cash_row.weight is not None else Decimal(0))
    if not any(m.code == "weight_sum" for m in msgs) and abs(total - 1) > tol:
        if total < 1 and fill_with_cash:
            gap = 1 - total
            if cash_row is None:
                last = max((r.row_num for r in rows), default=header_row)
                cash_row = Row(row_num=last + 1, raw={"cash": "filled"}, is_cash=True, weight=Decimal(0), name="Cash")
                rows.append(cash_row)
            cash_row.weight = (cash_row.weight or Decimal(0)) + gap
            msgs.append(Message("cash_filled", "info", f"the {pct_text(gap)} gap is filled with cash (row "
                                                       f"{cash_row.row_num})", cash_row.row_num, "weight"))
        else:
            msgs.append(Message("weight_sum", "error",
                                f"weights sum to {pct_text(total)}, not 100% (±{tolerance_pct}%)"
                                + (" — mark a cash row (blank weight takes the remainder) or choose 'fill with cash'"
                                   if total < 1 else ""), header_row, "weight"))
    if cash_row is not None and cash_row.weight is not None:
        cash_target = q(cash_row.weight)
        cash_row.state = "cash"
    for r in rows:
        if r.weight is None:
            continue
        lo, hi = r.drift_min, r.drift_max
        for f, b in (("drift_min", lo), ("drift_max", hi)):
            if b is not None and (b < 0 or b > 1):
                msgs.append(Message("band_range", "error", f"row {r.row_num}: {FIELDS[f][0]} {pct_text(b)} is outside "
                                                           "0–100%", r.row_num, f))
        if (lo is not None and lo > r.weight) or (hi is not None and hi < r.weight):
            msgs.append(Message("band_order", "error", f"row {r.row_num}: the band {pct_text(lo) if lo is not None else '—'}"
                                                       f" to {pct_text(hi) if hi is not None else '—'} does not contain "
                                                       f"the target {pct_text(r.weight)}", r.row_num, "drift_min"))
    return basis, cash_target, msgs


def _alt_type(ident: str) -> str:
    s = ident.strip().upper()
    if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}\d", s):
        return "isin"
    if re.fullmatch(r"[0-9A-Z]{8}\d", s) and any(ch.isdigit() for ch in s[:8]):
        return "cusip"
    return "ticker"


async def resolve_rows(conn, rows: list[Row]) -> list[Message]:
    """Resolve every non-cash row (and its alternates) through the security
    master. CUSIP/ISIN first; ticker/SEDOL only when neither is given, and only
    when unambiguous. Never guesses."""
    from services.securities_global import get_by_identifiers_bulk, normalize_identifier_value

    msgs: list[Message] = []
    want: dict[str, set] = {"cusip": set(), "isin": set(), "sedol": set(), "ticker": set()}
    for r in rows:
        if r.is_cash:
            continue
        for t in want:
            v = getattr(r, t)
            if v:
                want[t].add(v)
        for a in r.alternates:
            a = a.get("identifier") if isinstance(a, dict) else a
            if a:
                want[_alt_type(a)].add(a)
    hits: dict[str, dict] = {}
    for t, vals in want.items():
        if vals:
            hits[t] = await get_by_identifiers_bulk(conn, t, sorted(vals))

    def look(t: str, v: str | None):
        if not v:
            return None
        try:
            return hits.get(t, {}).get(normalize_identifier_value(t, v))
        except Exception:  # noqa: BLE001 — a malformed identifier is simply not a match
            return None

    for r in rows:
        if r.is_cash:
            r.state = "cash"
            continue
        if r.global_security_id and r.state == "resolved":
            pass          # already resolved (an edited draft row whose identifiers did not change)
        else:
            _resolve_one(r, look, msgs)
        alts = []
        for a in r.alternates:
            if isinstance(a, dict):
                a = a.get("identifier") or ""
            if not a:
                continue
            t = _alt_type(a)
            h = look(t, a)
            alts.append({"identifier": a, "id_type": t, "global_security_id": h["id"] if h else None})
            if not h:
                msgs.append(Message("alternate_unresolved", "info", f"row {r.row_num}: alternate {a} is not resolved",
                                    r.row_num, "alternates"))
        r.alternates = alts
    _check_duplicates(rows, msgs)
    return msgs


def _resolve_one(r: Row, look, msgs: list[Message]) -> None:
    strong = [(t, getattr(r, t), look(t, getattr(r, t))) for t in ("cusip", "isin") if getattr(r, t)]
    if strong:
        ids = {h["id"] for _, _, h in strong if h}
        missing = [f"{t.upper()} {v}" for t, v, h in strong if not h]
        if len(ids) > 1:
            r.state, r.resolution_note = "unresolved", "CUSIP and ISIN point at different securities"
            msgs.append(Message("identifier_conflict", "warning", f"row {r.row_num}: the CUSIP and ISIN point at "
                                                                  "different securities — queued for review", r.row_num))
        elif ids and not missing:
            hit = next(h for _, _, h in strong if h)
            r.global_security_id, r.security_name, r.state = hit["id"], hit.get("name"), "resolved"
        else:
            r.state = "unresolved"
            r.resolution_note = (f"{', '.join(missing)} not in the security master" if not ids
                                 else f"{', '.join(missing)} not in the security master (the other identifier "
                                      "resolves) — queued, not guessed")
            msgs.append(Message("unresolved", "warning", f"row {r.row_num}: {r.resolution_note}; it is queued and "
                                                         "never matched on its ticker", r.row_num,
                                "cusip" if r.cusip else "isin"))
    else:
        hit = look("sedol", r.sedol) or (None if r.sedol else look("ticker", r.ticker))
        if hit:
            r.global_security_id, r.security_name, r.state = hit["id"], hit.get("name"), "resolved"
        elif r.ticker or r.sedol:
            r.state = "unresolved"
            r.resolution_note = (f"{'SEDOL ' + r.sedol if r.sedol else 'ticker ' + r.ticker} is not in the security "
                                 "master or names more than one security")
            msgs.append(Message("unresolved", "warning", f"row {r.row_num}: {r.resolution_note} — queued", r.row_num,
                                "sedol" if r.sedol else "ticker"))
        else:
            r.state = "unresolved"
            msgs.append(Message("no_identifier", "error", f"row {r.row_num}: no ticker, CUSIP, ISIN or SEDOL",
                                r.row_num))


def _check_duplicates(rows: list[Row], msgs: list[Message]) -> None:
    seen_sec: dict[str, int] = {}
    seen_raw: dict[str, int] = {}
    for r in rows:
        if r.is_cash:
            continue
        if r.global_security_id:
            if r.global_security_id in seen_sec:
                msgs.append(Message("duplicate_security", "error", f"row {r.row_num}: the same security as row "
                                                                   f"{seen_sec[r.global_security_id]}", r.row_num))
            else:
                seen_sec[r.global_security_id] = r.row_num
        else:
            k = r.key()
            if k in seen_raw:
                msgs.append(Message("duplicate_identifier", "error", f"row {r.row_num}: the same identifier as row "
                                                                     f"{seen_raw[k]}", r.row_num))
            else:
                seen_raw[k] = r.row_num


async def parse_and_validate(conn, wb: Workbook, *, sheet: str | None, header_row: int | None,
                             mapping: dict | None, saved_mapping: dict | None, fill_with_cash: bool,
                             tolerance_pct: Decimal = DEFAULT_TOLERANCE_PCT) -> ParseResult:
    sheet = sheet or wb.sheets[0]
    if sheet not in wb.grids:
        raise UploadRefused("unreadable", f"the workbook has no sheet named {sheet!r}")
    grid = wb.grids[sheet]
    hdr = header_row or detect_header(grid, saved_mapping)
    msgs: list[Message] = []
    if hdr is None:
        return ParseResult(sheet, 0, [], {}, None, [], [Message(
            "no_header", "error", "no header row found in the first 30 rows (it must name an identifier column and a "
                                  "weight column) — choose the header row", None)], tolerance=tolerance_pct)
    headers = headers_of(grid, hdr)
    cols = mapping_from_request(headers, mapping, saved_mapping)
    fields = set(cols.values())
    if "weight" not in fields or not (fields & set(IDENTIFIER_FIELDS)):
        return ParseResult(sheet, hdr, headers, mapping_by_header(headers, cols), None, [], [Message(
            "mapping", "error", "map a weight column and at least one identifier column (Ticker, CUSIP, ISIN or SEDOL)",
            hdr)], tolerance=tolerance_pct)
    rows, m1 = extract_rows(wb, sheet, hdr, cols)
    msgs += m1
    if not rows:
        msgs.append(Message("no_rows", "error", "the sheet has no position rows under the header", hdr))
    basis, cash_target, m2 = interpret_weights(rows, header_row=hdr, tolerance_pct=tolerance_pct,
                                               fill_with_cash=fill_with_cash)
    msgs += m2
    msgs += await resolve_rows(conn, rows)
    order = {"error": 0, "warning": 1, "info": 2}
    msgs.sort(key=lambda m: (m.row if m.row is not None else 0, order[m.severity]))
    return ParseResult(sheet, hdr, headers, mapping_by_header(headers, cols), basis, rows, msgs, cash_target,
                       tolerance_pct)


async def validate_positions(conn, positions: list[dict], *, tolerance_pct: Decimal,
                             require_sum: bool = True) -> tuple[list[Row], list[Message]]:
    """The upload validation, applied to positions edited in the UI (weights
    given as fractions). Same resolution, duplicate, band and sum rules."""
    rows = []
    for p in positions:
        r = Row(row_num=int(p["row_num"]), raw={}, ticker=(p.get("raw_ticker") or "").upper() or None,
                cusip=(p.get("raw_cusip") or "").upper() or None, isin=(p.get("raw_isin") or "").upper() or None,
                sedol=(p.get("raw_sedol") or "").upper() or None, name=p.get("raw_name"),
                weight=p.get("target_weight"), asset_class=p.get("asset_class"), sleeve=p.get("sleeve_label"),
                drift_min=p.get("drift_min"), drift_max=p.get("drift_max"),
                alternates=[a["identifier"] if isinstance(a, dict) else str(a) for a in (p.get("alternates") or [])],
                restrictions=p.get("restrictions"), notes=p.get("notes"),
                is_cash=p.get("resolution_state") == "cash")
        if p.get("global_security_id") and p.get("resolution_state") == "resolved":
            r.global_security_id, r.state = str(p["global_security_id"]), "resolved"
        rows.append(r)
    msgs: list[Message] = []
    for r in rows:
        if r.weight is None:
            msgs.append(Message("weight_missing", "error", f"row {r.row_num}: no weight", r.row_num, "weight"))
        elif r.weight < 0:
            msgs.append(Message("weight_negative", "error", f"row {r.row_num}: negative weight", r.row_num, "weight"))
        elif r.weight > 1:
            msgs.append(Message("weight_range", "error", f"row {r.row_num}: weight above 100%", r.row_num, "weight"))
        for f in ("drift_min", "drift_max"):
            b = getattr(r, f)
            if b is not None and (b < 0 or b > 1):
                msgs.append(Message("band_range", "error", f"row {r.row_num}: {FIELDS[f][0]} outside 0–100%", r.row_num, f))
        if r.weight is not None and ((r.drift_min is not None and r.drift_min > r.weight)
                                     or (r.drift_max is not None and r.drift_max < r.weight)):
            msgs.append(Message("band_order", "error", f"row {r.row_num}: the band does not contain the target",
                                r.row_num, "drift_min"))
    if sum(1 for r in rows if r.is_cash) > 1:
        msgs.append(Message("cash_multiple", "error", "more than one cash row", None))
    total = sum((r.weight for r in rows if r.weight is not None), Decimal(0))
    if require_sum and abs(total - 1) > tolerance_pct / 100:
        msgs.append(Message("weight_sum", "error", f"weights sum to {pct_text(total)}, not 100% (±{tolerance_pct}%)",
                            None, "weight"))
    msgs += await resolve_rows(conn, rows)
    return rows, msgs


# ── tradability (informational, never blocking) ─────────────────────────────
TRADABILITY_KEYS = ("fractional_shares", "trade_list_upload", "model_api")


async def tradability(conn, org_id) -> list[dict]:
    from services.custodian_capabilities import available, capability

    codes = [r["c"] for r in await conn.fetch(
        "SELECT DISTINCT lower(custodian_code) AS c FROM public.accounts WHERE org_id = $1::uuid AND valid_to IS NULL "
        "AND system_to IS NULL AND NOT is_held_away ORDER BY 1", str(org_id))]
    if not codes:
        return [{"code": "tradability", "severity": "info", "custodian": None, "capability": None, "value": None,
                 "text": "no custodian accounts in this organization yet — tradability not assessed"}]
    out = []
    for code in codes:
        for key in TRADABILITY_KEYS:
            cap = await capability(conn, code, key)
            if not available(cap):
                out.append({"code": "tradability", "severity": "info", "custodian": code, "capability": key,
                            "value": cap.value, "text": f"{key} {cap.value} at {code}"})
    return out


# ── diff ────────────────────────────────────────────────────────────────────
def diff_positions(old: dict[str, dict], new: dict[str, dict], threshold: Decimal = DIFF_THRESHOLD) -> dict:
    """``old`` / ``new``: {key: {"weight": Decimal, "label": str}}. Exact
    Decimal: added, removed, weight moves over the threshold, and turnover =
    half the sum of every absolute weight change (including small ones)."""
    added, removed, changed = [], [], []
    turnover = Decimal(0)
    for k in sorted(set(old) | set(new)):
        o = old.get(k, {}).get("weight") if k in old else None
        n = new.get(k, {}).get("weight") if k in new else None
        label = (new.get(k) or old.get(k) or {}).get("label")
        if o is None:
            added.append({"key": k, "label": label, "weight": str(n)})
            turnover += abs(n)
        elif n is None:
            removed.append({"key": k, "label": label, "weight": str(o)})
            turnover += abs(o)
        else:
            d = n - o
            turnover += abs(d)
            if abs(d) > threshold:
                changed.append({"key": k, "label": label, "from": str(o), "to": str(n), "delta": str(d)})
    return {"added": added, "removed": removed, "changed": changed, "threshold": str(threshold),
            "turnover": str(turnover / 2), "has_base": bool(old)}


# ── template ────────────────────────────────────────────────────────────────
TEMPLATE_FIELDS = ("ticker", "cusip", "isin", "sedol", "name", "weight", "asset_class", "sleeve", "drift_min",
                   "drift_max", "alternates", "restrictions", "notes", "cash")


def template_headers() -> list[str]:
    return [FIELDS[f][0] for f in TEMPLATE_FIELDS]


def template_csv() -> bytes:
    buf = io.StringIO()
    csv.writer(buf, lineterminator="\r\n").writerow(template_headers())
    return buf.getvalue().encode("utf-8")


def template_xlsx() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws.append(template_headers())
    info = wb.create_sheet("Instructions")
    for line in (
        ["One row per position. Give a CUSIP or ISIN where you can — a ticker is used only when neither is given,"],
        ["and only when it names exactly one security. Unresolved rows are kept and queued, never guessed."],
        ["Target Weight: percentages (summing to 100) or fractions (summing to 1) — either is detected."],
        ["Min / Max: the drift band, in the same unit as the weight. Leave blank for none."],
        ["Cash: mark the cash row (yes); a cash row with a blank weight takes the remainder."],
        ["Alternate: other identifiers separated by semicolons. Values only — formulas are not evaluated."],
    ):
        info.append(line)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def vocabulary() -> dict:
    return {
        "mapping_fields": [{"key": k, "label": v[0], "synonyms": list(v[1])} for k, v in FIELDS.items()],
        "validation_codes": [{"key": k, "label": v} for k, v in VALIDATION_CODES.items()],
        "severities": [{"key": k, "label": v} for k, v in SEVERITY_LABELS.items()],
        "weight_bases": [{"key": "percent", "label": "Percent (sums to 100)"},
                         {"key": "fraction", "label": "Fraction (sums to 1)"}],
        "limits": {"max_bytes": MAX_BYTES, "max_rows": MAX_DATA_ROWS, "accepted": list(ACCEPTED_EXTENSIONS),
                   "refused": list(MACRO_EXTENSIONS)},
    }
