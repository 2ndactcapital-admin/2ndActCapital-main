"""verify_rep3 — Replication Studio, sprint 3 (REP3): model management — one model
registry, Excel/CSV upload with column mapping, a permanent upload log, the model
library (platform + org shelves, copy or subscribe), composite (sleeve) models,
approvals on the governed human-task path, and rep1's terms approvals moved onto it.

WRITTEN, NOT RUN by the sprint. The operator runs (after `pip install -r
apps/api/requirements.txt`, which adds xlrd for .xls):

    doppler run -- apps/api/venv/bin/python apps/api/scripts/verify_rep3.py 2>&1 | grep -E "^\\[FAIL\\]|^TOTAL"

Pass/fail only, no prompts. Prints 'TOTAL: N PASS, M FAIL' and exits non-zero
on any failure. AUTHZ_TEST_MODE=1 is set before anything imports ``main``;
secrets are hydrated from Doppler over HTTPS (``_db_connect``) first.

MOCKED: the token boundary (``main.verify_token``), the Auth0 profile call and
R2 (``services.storage.upload_bytes`` / ``download_bytes`` → an in-memory dict
that counts objects). NO network: the real httpx network transport is replaced
by one that records and refuses — any attempt is a failed check.

REAL: the app (middleware, Principal, every gate, every handler) through an
ASGI transport, the workflow engine, RLS, the live database. The script's own
connection is DATABASE_URL = app_service (rolbypassrls = false, asserted).

FIXTURES: three orgs (99000000-…-0000000e3001 / e3002 / e3003) and ONE fixture
user in the real Hollisworks platform org (super admin, is_test) who publishes
the platform library entry. Users are created explicitly with the auth0_sub
their token carries and an @test.local email (is_test by trigger). Fixture
securities carry the reserved prefix ZZRM (asserted absent from the real
master first). The only rows written in the platform org are the fixture
platform model (created_by = the fixture user), its versions and positions, the
library entry, and the publish event (+ its audit and domain event) — all
deleted by those ids.

TEARDOWN (start and end): the operator DSN (role postgres) with
models.teardown / replica.teardown = 'fixture' in its transaction — the only
path past the immutability triggers — model tables in their FK order (the
models ↔ library-entries cycle broken by deleting library-sourced models
first), the rest child-before-parent from pg_constraint, every DELETE keyed by
fixture ids; then every counted table is re-read for its exact pre-test count
and real models / securities / fee schedules / workflow rows / users are
compared by count + checksum.

With VERIFY_NESTED=1 the regression section prints one '[SKIP] nested run'.
"""

from __future__ import annotations

from pathlib import Path
import os as _authz_os
_authz_os.environ["AUTHZ_TEST_MODE"] = "1"  # AUTHZ1: set before main is imported

import asyncio
import csv
import io
import json
import os
import pathlib
import re
import struct
import subprocess
import sys
import traceback
import uuid
from datetime import date, datetime, timezone
from decimal import Decimal
from uuid import UUID

HERE = pathlib.Path(__file__).resolve().parent
API_DIR = HERE.parent
REPO = API_DIR.parents[1]
WEB_DIR = REPO / "apps" / "web"
sys.path.insert(0, str(HERE))

from _db_connect import admin_dsn, connect, hydrate, resolve_dsn  # noqa: E402

hydrate()  # BEFORE main is imported (Settings is lru_cached)

import asyncpg  # noqa: E402
import httpx  # noqa: E402
import openpyxl  # noqa: E402

import _authz1_common as C  # noqa: E402

BEFORE_REF = "93aaf23"   # rep3-work base — main after "merge rep2" + schema refresh
TAG = "rep3verify"
PREFIX = "ZZRM"
OPERATOR_DSN_VAR = "DATABASE_URL_PRECUTOVER_POSTGRES_BACKUP"
PLATFORM_ORG = UUID("bb347258-8f28-4f49-8cc9-e29ccad82884")
SEC_NAMES = ("REP3 VERIFY ALPHA", "REP3 VERIFY BETA", "REP3 VERIFY GAMMA", "REP3 VERIFY DECOY", "REP3 VERIFY FUND")


def _id(n: int) -> UUID:
    return UUID(f"99000000-0000-0000-0000-0000000e3{n:03d}")


ORG_A, ORG_B, ORG_C = _id(1), _id(2), _id(3)
FIXTURE_ORGS = [ORG_A, ORG_B, ORG_C]
(U_SUPER, U_ADMIN1, U_ADMIN2, U_VIEWER, U_MEMBER, U_ADMIN_B1, U_ADMIN_B2, U_NEWADMIN, U_PLAT) = (
    _id(110 + i) for i in range(9))
USERS = {   # org, users.role, short, rbac role, full name
    U_SUPER: (ORG_A, "super_admin", "super", "super_admin", "REP3 Super Admin"),
    U_ADMIN1: (ORG_A, "org_admin", "admin1", "org_admin", "REP3 Admin Maker"),
    U_ADMIN2: (ORG_A, "org_admin", "admin2", "org_admin", "REP3 Admin Checker"),
    U_VIEWER: (ORG_A, "member", "viewer", "rep3_viewer", "REP3 Viewer"),
    U_MEMBER: (ORG_A, "member", "member", "member", "REP3 Member"),
    U_ADMIN_B1: (ORG_B, "org_admin", "adminb1", "org_admin", "REP3 Org B Admin"),
    U_ADMIN_B2: (ORG_B, "org_admin", "adminb2", "org_admin", "REP3 Org B Checker"),
    U_NEWADMIN: (ORG_C, "member", "newadmin", None, "REP3 New Org Admin"),
    U_PLAT: (PLATFORM_ORG, "super_admin", "plat", None, "REP3 Platform Publisher"),
}
MODEL_PERMS = ("view_models", "manage_models", "approve_models")
ENT_A, ACCT_A = _id(201), _id(301)
CODE_ALT = "ALTRUIST"
CUSIP = {"A": "ZZRM00011", "B": "ZZRM00012", "C": "ZZRM00013", "D": "ZZRM00019", "F": "ZZRM00021"}
TICKER = {"A": "ZZRMA", "B": "ZZRMB", "C": "ZZRMC", "D": "ZZRMX", "F": "ZZRMF"}
MISSING_CUSIP = "ZZRM00099"      # in no master row; its row carries the DECOY's ticker

COUNTED = (
    "public.users", "public.organizations", "public.roles", "public.user_roles", "public.role_permissions",
    "public.audit_log", "public.member_todos", "public.domain_events", "public.domain_event_deliveries",
    "public.workflow_definitions", "public.workflow_versions", "public.workflow_steps", "public.workflow_runs",
    "public.workflow_run_steps", "public.entities", "public.accounts", "public.org_settings",
    "public.document_field_corrections",
    "portfolio.securities_global", "portfolio.securities_global_identifiers", "portfolio.models",
    "portfolio.model_versions", "portfolio.model_version_positions", "portfolio.model_sleeves",
    "portfolio.model_uploads", "portfolio.model_upload_events", "portfolio.model_column_mappings",
    "portfolio.model_library_entries", "portfolio.model_subscriptions", "portfolio.model_events",
    "portfolio.replica_securities", "portfolio.replica_terms_versions", "portfolio.replica_terms_actions",
)

S: dict = {}
NETWORK_CALLS: list[str] = []
STORE: dict[str, bytes] = {}
PUTS: list[str] = []
_n_pass = 0
_n_fail = 0


def check(passed: bool, label: str, detail: str = "") -> bool:
    assert isinstance(passed, bool), f"check() got {passed!r} ({type(passed).__name__}) for {label!r}"
    global _n_pass, _n_fail
    print(f"{'[PASS]' if passed else '[FAIL]'} {label}" + (f"  — {detail}" if detail else ""))
    if passed:
        _n_pass += 1
    else:
        _n_fail += 1
    return passed


def find(label: str, detail: str = "") -> None:
    print(f"[FIND] {label}" + (f"  — {detail}" if detail else ""))


def section(title: str) -> None:
    print(f"\n── {title} ──")


def run_cmd(args, cwd, timeout=600, env=None):
    try:
        p = subprocess.run(args, cwd=cwd, capture_output=True, text=True, timeout=timeout, env=env)
        return p.returncode, p.stdout, p.stderr
    except subprocess.TimeoutExpired as exc:
        return 124, exc.stdout or "", f"timeout after {timeout}s"


def git_show(ref: str, path: str) -> str:
    code, out, _ = run_cmd(["git", "show", f"{ref}:{path}"], REPO, timeout=60)
    return out if code == 0 else ""


def sub_of(uid: UUID) -> str:
    return f"auth0|{TAG}-{USERS[uid][2]}"


def email_of(uid: UUID) -> str:
    return f"{TAG}-{USERS[uid][2]}@test.local"


def D(v) -> Decimal | None:
    return None if v is None else Decimal(str(v))


# ═══ DB helpers (verify's own app_service connection; RLS context explicit) ═══
async def platform(conn, sql, *args, fetch="fetch"):
    async with conn.transaction():
        await conn.execute("SELECT set_config('app.is_super_admin', 'true', true)")
        return await getattr(conn, fetch)(sql, *args)


async def scoped(conn, org, sql, *args, fetch="fetch"):
    async with conn.transaction():
        await conn.execute(
            "SELECT set_config('app.current_org_id', $1, true), set_config('app.is_super_admin', 'false', true),"
            " set_config('app.current_auth0_sub', '', true)", str(org))
        return await getattr(conn, fetch)(sql, *args)


async def counts(conn) -> dict:
    return {t: int(await platform(conn, f"SELECT count(*) FROM {t}", fetch="fetchval")) for t in COUNTED}


async def fingerprint(conn, sql, *args) -> tuple[int, str]:
    row = await platform(conn, f"SELECT count(*) AS n, md5(COALESCE(string_agg(t::text, '' ORDER BY t::text), '')) "
                               f"AS s FROM ({sql}) t", *args, fetch="fetchrow")
    return int(row["n"]), row["s"]


async def real_fingerprints(conn) -> dict:
    """Real rows this sprint must never touch (count + checksum)."""
    orgs = [str(o) for o in FIXTURE_ORGS]
    users = [str(u) for u in USERS]
    return {
        "models": await fingerprint(conn, "SELECT * FROM portfolio.models WHERE NOT (org_id = ANY($1::uuid[])) "
                                          "AND (created_by IS NULL OR NOT (created_by = ANY($2::uuid[])))", orgs, users),
        "model_library_entries": await fingerprint(conn, "SELECT * FROM portfolio.model_library_entries WHERE NOT "
                                                         "(published_by = ANY($1::uuid[]))", users),
        "fee_schedules": await fingerprint(conn, "SELECT * FROM public.fee_schedules"),
        "securities_global": await fingerprint(conn, "SELECT * FROM portfolio.securities_global WHERE NOT "
                                                     "(name = ANY($1::text[]))", list(SEC_NAMES)),
        "users": await fingerprint(conn, "SELECT id, org_id, role, status, email FROM public.users WHERE NOT "
                                         "(id = ANY($1::uuid[]))", users),
        "workflow_definitions": await fingerprint(
            conn, "SELECT * FROM public.workflow_definitions WHERE NOT (org_id = ANY($1::uuid[]))", orgs),
        "workflow_runs": await fingerprint(
            conn, "SELECT * FROM public.workflow_runs WHERE NOT (org_id = ANY($1::uuid[]))", orgs),
        "workflow_run_steps": await fingerprint(
            conn, "SELECT * FROM public.workflow_run_steps WHERE NOT (org_id = ANY($1::uuid[]))", orgs),
        "replica_terms_versions": await fingerprint(
            conn, "SELECT * FROM portfolio.replica_terms_versions WHERE NOT (org_id = ANY($1::uuid[]))", orgs),
    }


async def operator_conn():
    dsn, how = await resolve_dsn(OPERATOR_DSN_VAR)
    if not dsn:
        raise RuntimeError(f"no operator DSN for the fixture teardown path: {how}")
    return await connect(dsn)


_FK_CACHE: dict = {}


async def fk_children(conn, table: str):
    if table not in _FK_CACHE:
        _FK_CACHE[table] = await conn.fetch("""
            SELECT cn.nspname || '.' || cc.relname AS child,
                   array_agg(ca.attname ORDER BY k.ord) AS child_cols,
                   array_agg(pa.attname ORDER BY k.ord) AS parent_cols
            FROM pg_constraint c
            JOIN LATERAL unnest(c.conkey, c.confkey) WITH ORDINALITY AS k(ck, pk, ord) ON true
            JOIN pg_class cc ON cc.oid = c.conrelid JOIN pg_namespace cn ON cn.oid = cc.relnamespace
            JOIN pg_attribute ca ON ca.attrelid = c.conrelid AND ca.attnum = k.ck
            JOIN pg_attribute pa ON pa.attrelid = c.confrelid AND pa.attnum = k.pk
            WHERE c.contype::text = 'f' AND c.confrelid = $1::regclass
            GROUP BY c.oid, cn.nspname, cc.relname""", table)
    return _FK_CACHE[table]


async def delete_tree(conn, table: str, where: str, params: list, depth=0, log=None, path=()):
    """Child-before-parent delete for a FIXTURE-ID parent (users / roles / orgs
    of this run only), FKs read from pg_constraint (recursively)."""
    if depth > 14:
        raise RuntimeError(f"FK recursion too deep at {table} via {' > '.join(path)}")
    path = path + (table,)
    for fk in await fk_children(conn, table):
        if fk["child"] in path:
            continue
        cc = ", ".join(f'"{c}"' for c in fk["child_cols"])
        pc = ", ".join(f'"{c}"' for c in fk["parent_cols"])
        await delete_tree(conn, fk["child"], f"({cc}) IN (SELECT {pc} FROM {table} WHERE {where})",
                          params, depth + 1, log, path)
    status = await conn.execute(f"DELETE FROM {table} WHERE {where}", *params)
    n = int(status.rsplit(" ", 1)[-1])
    if n and log is not None:
        log.append(f"{table}: {n}")


async def fk_order(conn, tables: list[str]) -> list[str]:
    """``tables`` sorted child-before-parent from pg_constraint (Kahn)."""
    edges = {t: set() for t in tables}
    rows = await conn.fetch("""
        SELECT cn.nspname || '.' || cc.relname AS child, pn.nspname || '.' || pc.relname AS parent
        FROM pg_constraint c JOIN pg_class cc ON cc.oid = c.conrelid JOIN pg_namespace cn ON cn.oid = cc.relnamespace
        JOIN pg_class pc ON pc.oid = c.confrelid JOIN pg_namespace pn ON pn.oid = pc.relnamespace
        WHERE c.contype = 'f'""")
    for r in rows:
        if r["child"] in edges and r["parent"] in edges and r["child"] != r["parent"]:
            edges[r["child"]].add(r["parent"])
    children_of = {t: {c for c, ps in edges.items() if t in ps} for t in tables}
    out, done = [], set()
    while len(out) < len(tables):
        ready = [t for t in tables if t not in done and children_of[t] <= done]
        if not ready:
            raise RuntimeError(f"FK cycle among {sorted(set(tables) - done)}")
        for t in ready:
            out.append(t)
            done.add(t)
    return out


async def fixture_security_ids(conn) -> list:
    rows = await conn.fetch(
        "SELECT DISTINCT global_security_id FROM portfolio.securities_global_identifiers WHERE id_value LIKE $1",
        f"{PREFIX}%")
    return [r["global_security_id"] for r in rows]


async def teardown(_conn=None) -> list[str]:
    """Operator connection, models/replica teardown = 'fixture', every DELETE keyed by fixture ids."""
    log: list[str] = []
    op = await operator_conn()
    try:
        async with op.transaction():
            await op.execute("SELECT set_config('models.teardown', 'fixture', true), "
                             "set_config('replica.teardown', 'fixture', true), "
                             "set_config('execution.teardown', 'fixture', true)")
            secs = await fixture_security_ids(op)
            users = list(USERS)
            orgs = FIXTURE_ORGS
            plat_models = [r["id"] for r in await op.fetch(
                "SELECT id FROM portfolio.models WHERE org_id = $1 AND (created_by = ANY($2::uuid[]) OR "
                "archived_from_id IN (SELECT id FROM portfolio.models WHERE org_id = $1 AND created_by = ANY($2::uuid[])))",
                PLATFORM_ORG, users)]
            models = [r["id"] for r in await op.fetch(
                "SELECT id FROM portfolio.models WHERE org_id = ANY($1::uuid[])", orgs)] + plat_models
            versions = [r["id"] for r in await op.fetch(
                "SELECT id FROM portfolio.model_versions WHERE model_id = ANY($1::uuid[]) OR org_id = ANY($2::uuid[])",
                models, orgs)]
            entries = [r["id"] for r in await op.fetch(
                "SELECT id FROM portfolio.model_library_entries WHERE model_id = ANY($1::uuid[]) OR org_id = "
                "ANY($2::uuid[]) OR published_by = ANY($3::uuid[])", models, orgs, users)]
            events = [r["id"] for r in await op.fetch(
                "SELECT id FROM portfolio.model_events WHERE org_id = ANY($1::uuid[]) OR model_id = ANY($2::uuid[]) "
                "OR actor_id = ANY($3::uuid[]) OR library_entry_id = ANY($4::uuid[])", orgs, models, users, entries)]
            sources = events + versions
            steps = (
                ("model_events", "DELETE FROM portfolio.model_events WHERE id = ANY($1::uuid[])", [events]),
                ("model_upload_events", "DELETE FROM portfolio.model_upload_events WHERE org_id = ANY($1::uuid[])",
                 [orgs]),
                ("model_subscriptions", "DELETE FROM portfolio.model_subscriptions WHERE org_id = ANY($1::uuid[])",
                 [orgs]),
                ("model_sleeves", "DELETE FROM portfolio.model_sleeves WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("model_version_positions", "DELETE FROM portfolio.model_version_positions WHERE version_id = "
                                            "ANY($1::uuid[])", [versions]),
                ("model_versions", "DELETE FROM portfolio.model_versions WHERE id = ANY($1::uuid[])", [versions]),
                ("model_uploads", "DELETE FROM portfolio.model_uploads WHERE org_id = ANY($1::uuid[])", [orgs]),
                # the models <-> library entries cycle: library-sourced models first (no entry is built on one)
                ("models (library-sourced)", "DELETE FROM portfolio.models WHERE id = ANY($1::uuid[]) AND "
                                             "library_entry_id IS NOT NULL", [models]),
                ("model_library_entries", "DELETE FROM portfolio.model_library_entries WHERE id = ANY($1::uuid[])",
                 [entries]),
                ("models", "DELETE FROM portfolio.models WHERE id = ANY($1::uuid[])", [models]),
                ("model_column_mappings", "DELETE FROM portfolio.model_column_mappings WHERE org_id = ANY($1::uuid[])",
                 [orgs]),
                ("member_todos", "DELETE FROM member_todos WHERE org_id = ANY($1::uuid[]) OR user_id = ANY($2::uuid[])",
                 [orgs, users]),
                ("domain_event_deliveries", "DELETE FROM domain_event_deliveries WHERE org_id = ANY($1::uuid[]) OR "
                 "domain_event_id IN (SELECT id FROM domain_events WHERE source_id = ANY($2::uuid[]))", [orgs, sources]),
                ("domain_events", "DELETE FROM domain_events WHERE org_id = ANY($1::uuid[]) OR (org_id = $3 AND "
                                  "source_id = ANY($2::uuid[]))", [orgs, sources, PLATFORM_ORG]),
                ("audit_log", "DELETE FROM audit_log WHERE org_id = ANY($1::uuid[]) OR user_id = ANY($2::uuid[]) OR "
                              "resource_id = ANY($3::uuid[])", [orgs, users, sources]),
                ("workflow_run_steps", "DELETE FROM workflow_run_steps WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("workflow_runs", "DELETE FROM workflow_runs WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("workflow_steps", "DELETE FROM workflow_steps WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("workflow_versions", "DELETE FROM workflow_versions WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("workflow_definitions", "DELETE FROM workflow_definitions WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("replica_terms_actions", "DELETE FROM portfolio.replica_terms_actions WHERE org_id = ANY($1::uuid[])",
                 [orgs]),
                ("document_field_corrections", "DELETE FROM document_field_corrections WHERE org_id = ANY($1::uuid[])",
                 [orgs]),
                ("replica_terms_versions", "DELETE FROM portfolio.replica_terms_versions WHERE org_id = "
                                           "ANY($1::uuid[])", [orgs]),
                ("replica_securities", "DELETE FROM portfolio.replica_securities WHERE org_id = ANY($1::uuid[])",
                 [orgs]),
                ("accounts", "DELETE FROM public.accounts WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("entities", "DELETE FROM public.entities WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("org_settings", "DELETE FROM public.org_settings WHERE org_id = ANY($1::uuid[])", [orgs]),
                ("securities_global_identifiers", "DELETE FROM portfolio.securities_global_identifiers WHERE "
                                                  "global_security_id = ANY($1::uuid[])", [secs]),
                ("securities_global", "DELETE FROM portfolio.securities_global WHERE id = ANY($1::uuid[])", [secs]),
            )
            for label, sql, args in steps:
                status = await op.execute(sql, *args)
                n = int(status.rsplit(" ", 1)[-1])
                if n:
                    log.append(f"{label}: {n}")
            await delete_tree(op, "public.users", "id = ANY($1::uuid[]) OR auth0_sub LIKE $2",
                              [users, f"%{TAG}%"], log=log)
            await delete_tree(op, "public.roles", "org_id = ANY($1::uuid[])", [orgs], log=log)
            await delete_tree(op, "public.organizations", "id = ANY($1::uuid[])", [orgs], log=log)
    finally:
        await op.close()
    return log


async def leftovers(conn) -> int:
    n = 0
    orgs = FIXTURE_ORGS
    users = list(USERS)
    for sql, args in (
        ("SELECT count(*) FROM users WHERE id = ANY($1::uuid[]) OR auth0_sub LIKE $2", [users, f"%{TAG}%"]),
        ("SELECT count(*) FROM organizations WHERE id = ANY($1::uuid[])", [orgs]),
        ("SELECT count(*) FROM portfolio.securities_global_identifiers WHERE id_value LIKE $1", [f"{PREFIX}%"]),
        ("SELECT count(*) FROM portfolio.models WHERE org_id = ANY($1::uuid[]) OR created_by = ANY($2::uuid[])",
         [orgs, users]),
        ("SELECT count(*) FROM portfolio.model_library_entries WHERE published_by = ANY($1::uuid[])", [users]),
        ("SELECT count(*) FROM portfolio.model_events WHERE org_id = ANY($1::uuid[]) OR actor_id = ANY($2::uuid[])",
         [orgs, users]),
        *[(f"SELECT count(*) FROM {t} WHERE org_id = ANY($1::uuid[])", [orgs]) for t in (
            "portfolio.model_versions", "portfolio.model_version_positions", "portfolio.model_sleeves",
            "portfolio.model_uploads", "portfolio.model_upload_events", "portfolio.model_column_mappings",
            "portfolio.model_subscriptions", "public.workflow_definitions", "public.workflow_runs",
            "public.workflow_run_steps", "public.member_todos", "public.domain_events", "public.audit_log",
            "public.accounts", "public.entities", "portfolio.replica_securities", "portfolio.replica_terms_versions")],
        ("SELECT count(*) FROM audit_log WHERE user_id = ANY($1::uuid[])", [users]),
    ):
        n += int(await platform(conn, sql, *args, fetch="fetchval"))
    return n


# ═══ App access ═════════════════════════════════════════════════════════════
CLAIMS_BY_TOKEN: dict[str, dict] = {}


def token_for(uid: UUID) -> str:
    c = {"sub": sub_of(uid), "org_id": str(USERS[uid][0])}
    key = "t" + str(abs(hash(json.dumps(c, sort_keys=True))))
    CLAIMS_BY_TOKEN[key] = c
    return key


async def api_raw(uid, method, path, body=None, files=None, form=None, timeout=240.0):
    import main

    kw = {"headers": {"Authorization": f"Bearer {token_for(uid)}", "User-Agent": "verify_rep3"}}
    if body is not None:
        kw["json"] = body
    if files is not None:
        kw["files"] = files
        kw["data"] = form or {}
    async with httpx.AsyncClient(transport=httpx.ASGITransport(app=main.app, raise_app_exceptions=False),
                                 base_url="http://verify", timeout=timeout) as client:
        try:
            res = await asyncio.wait_for(client.request(method, path, **kw), timeout)
        except asyncio.TimeoutError:
            return 599, b"", {}
    return res.status_code, res.content, dict(res.headers)


async def api(uid, method, path, body=None, timeout=240.0):
    code, content, _ = await api_raw(uid, method, path, body, timeout=timeout)
    try:
        return code, json.loads(content.decode("utf-8") or "{}")
    except ValueError:
        return code, {"raw": content[:300].decode("utf-8", "replace")}


async def upload(uid, name: str, data: bytes, provider=None, model_id=None, ctype="application/octet-stream"):
    form = {}
    if provider:
        form["provider_name"] = provider
    if model_id:
        form["model_id"] = str(model_id)
    code, content, _ = await api_raw(uid, "POST", "/api/v1/models/uploads", files={"file": (name, data, ctype)},
                                     form=form)
    try:
        return code, json.loads(content.decode("utf-8") or "{}")
    except ValueError:
        return code, {"raw": content[:300].decode("utf-8", "replace")}


def install_stubs():
    import main
    from services import storage
    from services import users as users_mod

    main.verify_token = lambda tok: dict(CLAIMS_BY_TOKEN[tok])
    from services.brief_blocks import register_brief_blocks

    register_brief_blocks()

    async def no_profile(request, c):
        return None

    users_mod.fetch_auth0_profile = no_profile

    def put(key, data, content_type=None, bucket=None):
        PUTS.append(key)
        STORE[key] = bytes(data)
        return key

    def get(key, bucket=None):
        return STORE[key]

    storage.upload_bytes = put
    storage.download_bytes = get

    async def refuse_network(self, request):
        NETWORK_CALLS.append(str(request.url))
        raise AssertionError(f"verify_rep3: a network call was attempted: {request.url}")

    httpx.AsyncHTTPTransport.handle_async_request = refuse_network


class _ctx:
    """RLS contextvars for a direct service call (what the middleware sets)."""

    def __init__(self, org, super_admin=False):
        self.org, self.super_admin = org, super_admin

    def __enter__(self):
        from services.database import set_rls_context

        self.tokens = set_rls_context(str(self.org) if self.org else None, self.super_admin)
        return self

    def __exit__(self, *exc):
        from services.database import reset_rls_context

        reset_rls_context(self.tokens)
        return False


# ═══ Fixture files ══════════════════════════════════════════════════════════
HEADER = ["Ticker", "CUSIP", "Weight %", "Min", "Max", "Asset Class"]
ROWS_PCT = [[TICKER["A"], CUSIP["A"], 50, 40, 60, "Equity"],
            [TICKER["B"], CUSIP["B"], 30, None, None, "Equity"],
            [TICKER["C"], None, 20, None, None, "Bond"]]


def csv_bytes(header, rows, bom=True, delimiter=",") -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delimiter, lineterminator="\r\n")
    w.writerow(header)
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return (("﻿" if bom else "") + buf.getvalue()).encode("utf-8")


def xlsx_bytes(header, rows, extra_sheet=True) -> bytes:
    wb = openpyxl.Workbook()
    if extra_sheet:
        wb.active.title = "Cover"
        wb.active.append(["Model file — see the Holdings sheet"])
        ws = wb.create_sheet("Holdings")
    else:
        ws = wb.active
    ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _rec(rtype: int, payload: bytes) -> bytes:
    return struct.pack("<HH", rtype, len(payload)) + payload


def _ustr8(s: str) -> bytes:        # 1-byte length, compressed (latin-1) characters
    return struct.pack("<BB", len(s), 0) + s.encode("latin-1")


def xls_bytes(header, rows, sheet="Holdings") -> bytes:
    """A real BIFF8 workbook in an OLE2 compound document, written here (no xlwt):
    globals (BOF, CODEPAGE, FONT, two XFs, BOUNDSHEET, SST, EOF) + one worksheet
    (BOF, DIMENSIONS, LABELSST / NUMBER cells, EOF). Numbers are NUMBER records
    (IEEE double); text goes through the shared-string table."""
    grid = [header] + rows
    strings: list[str] = []
    for r in grid:
        for v in r:
            if isinstance(v, str) and v not in strings:
                strings.append(v)
    bof = lambda kind: _rec(0x0809, struct.pack("<HHHHII", 0x0600, kind, 0x0DBB, 0x07CC, 0, 0x06))  # noqa: E731
    font = _rec(0x0031, struct.pack("<HHHHHBBBB", 200, 0, 0x7FFF, 400, 0, 0, 0, 0, 0) + _ustr8("Arial"))
    xf_style = _rec(0x00E0, struct.pack("<HHHBBBBIiH", 0, 0, 0xFFF5, 0x20, 0, 0, 0, 0, 0, 0x20C0))
    xf_cell = _rec(0x00E0, struct.pack("<HHHBBBBIiH", 0, 0, 0x0001, 0x20, 0, 0, 0, 0, 0, 0x20C0))
    sst_body = struct.pack("<II", sum(1 for r in grid for v in r if isinstance(v, str)), len(strings))
    for s in strings:
        sst_body += struct.pack("<HB", len(s), 0) + s.encode("latin-1")
    sst = _rec(0x00FC, sst_body)
    head = bof(0x0005) + _rec(0x0042, struct.pack("<H", 1200)) + font + xf_style + xf_cell
    boundsheet_len = 4 + 6 + 2 + len(sheet)          # record header + pos/vis/type + ustr8
    tail = sst + _rec(0x000A, b"")
    sheet_pos = len(head) + boundsheet_len + len(tail)
    boundsheet = _rec(0x0085, struct.pack("<IBB", sheet_pos, 0, 0) + _ustr8(sheet))
    globals_ = head + boundsheet + tail
    assert len(globals_) == sheet_pos
    ncols = max(len(r) for r in grid)
    cells = b""
    for ri, r in enumerate(grid):
        for ci, v in enumerate(r):
            if v is None:
                continue
            if isinstance(v, str):
                cells += _rec(0x00FD, struct.pack("<HHHI", ri, ci, 1, strings.index(v)))
            else:
                cells += _rec(0x0203, struct.pack("<HHHd", ri, ci, 1, float(v)))
    ws = bof(0x0010) + _rec(0x0200, struct.pack("<IIHHH", 0, len(grid), 0, ncols, 0)) + cells + _rec(0x000A, b"")
    stream = globals_ + ws
    stream += b"\x00" * max(0, 4096 - len(stream))       # >= the mini-stream cutoff: a regular stream
    return _ole2(stream)


def _ole2(stream: bytes) -> bytes:
    sec = 512
    n_stream = -(-len(stream) // sec)
    data = stream + b"\x00" * (n_stream * sec - len(stream))
    endchain, freesect, fatsect, nostream = 0xFFFFFFFE, 0xFFFFFFFF, 0xFFFFFFFD, 0xFFFFFFFF
    dir_sector = 1 + n_stream
    fat = [fatsect] + [i + 2 if i < n_stream - 1 else endchain for i in range(n_stream)] + [endchain]
    fat += [freesect] * (sec // 4 - len(fat))
    assert len(fat) == sec // 4

    def dirent(name, etype, child, start, size):
        nm = (name.encode("utf-16-le") + b"\x00\x00") if name else b""
        return (nm.ljust(64, b"\x00") + struct.pack("<HBB", len(nm), etype, 1)
                + struct.pack("<III", nostream, nostream, child) + b"\x00" * 16 + struct.pack("<I", 0)
                + b"\x00" * 16 + struct.pack("<III", start, size, 0))

    dirs = (dirent("Root Entry", 5, 1, endchain, 0) + dirent("Workbook", 2, nostream, 1, len(stream))
            + dirent("", 0, nostream, freesect, 0) * 2)
    assert len(dirs) == sec
    header = (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 16
              + struct.pack("<HHHHH", 0x003E, 0x0003, 0xFFFE, 9, 6) + b"\x00" * 6
              + struct.pack("<IIIIIIIII", 0, 1, dir_sector, 0, 4096, endchain, 0, endchain, 0)
              + struct.pack("<I", 0) + struct.pack("<I", freesect) * 108)
    assert len(header) == sec
    return header + struct.pack(f"<{sec // 4}I", *fat) + data + dirs


# ═══ Seed ═══════════════════════════════════════════════════════════════════
async def seed(conn) -> None:
    from services import rbac
    from services import securities_global as SG
    from services.custody.base import AccountNumber
    from services.custody.registry import get_or_create_salt

    async with conn.transaction():
        await conn.execute("SELECT set_config('app.is_super_admin', 'true', true)")
        for org, nm in ((ORG_A, "a"), (ORG_B, "b"), (ORG_C, "c")):
            await conn.execute("INSERT INTO organizations (id, name, slug) VALUES ($1, $2, $3)",
                               org, f"REP3 Verify Org {nm.upper()}", f"{TAG}-{nm}")
        perm_ids = {r["name"]: r["id"] for r in await conn.fetch("SELECT id, name FROM permissions")}
        roles = {}
        for org, rname in ((ORG_A, "member"), (ORG_A, "super_admin")):
            desc, perms = C.BASELINE_ROLES[rname]
            rid = await rbac.ensure_role(conn, org, rname, desc)
            for p in perms:
                if p in perm_ids:
                    await conn.execute("INSERT INTO role_permissions (role_id, permission_id) VALUES ($1, $2) "
                                       "ON CONFLICT DO NOTHING", rid, perm_ids[p])
            roles[(org, rname)] = rid
        rid = await rbac.ensure_role(conn, ORG_A, "rep3_viewer", "REP3 verify — view_models only")
        await conn.execute("INSERT INTO role_permissions (role_id, permission_id) VALUES ($1, $2)", rid,
                           perm_ids["view_models"])
        roles[(ORG_A, "rep3_viewer")] = rid
        for uid, (org, account_role, nm, rname, full) in USERS.items():
            await conn.execute(
                "INSERT INTO users (id, org_id, email, full_name, auth0_sub, role, status, is_active) "
                "VALUES ($1, $2, $3, $4, $5, $6, 'active', true)",
                uid, org, email_of(uid), full, sub_of(uid), account_role)
            if rname == "org_admin":
                await rbac.grant_org_admin(conn, uid, org)
            elif rname:
                await conn.execute("INSERT INTO user_roles (user_id, role_id) VALUES ($1, $2)", uid,
                                   roles[(org, rname)])
        await conn.execute("INSERT INTO entities (id, org_id, entity_type, display_name) VALUES ($1, $2, "
                           "'individual', 'REP3 Verify Client')", ENT_A, ORG_A)
        for k, name in zip("ABCDF", SEC_NAMES):
            sid = await SG.create_security(conn, name=name, security_type="fund" if k == "F" else "equity",
                                           is_super_admin=True)
            await SG.add_identifier(conn, global_security_id=sid, id_type="cusip", id_value=CUSIP[k],
                                    is_super_admin=True)
            await SG.add_identifier(conn, global_security_id=sid, id_type="ticker", id_value=TICKER[k],
                                    is_super_admin=True)
            S[f"sec_{k}"] = str(sid)
        await conn.execute("SELECT set_config('app.current_org_id', $1, true)", str(ORG_A))
        salt = await get_or_create_salt(conn, ORG_A)
        n = AccountNumber("ZZ31415926")
        await conn.execute(
            "INSERT INTO accounts (id, org_id, account_number_masked, account_number_hash, custodian_code, "
            "registration_type, tax_status, primary_entity_id, is_held_away) VALUES ($1, $2, $3, $4, $5, "
            "'individual', 'taxable', $6, false)", ACCT_A, ORG_A, n.masked, n.hashed(salt), CODE_ALT, ENT_A)
    users = await platform(conn, "SELECT id FROM users WHERE id = ANY($1::uuid[]) AND is_test", list(USERS))
    check(len(users) == len(USERS), "fixture users created with the token's auth0_sub, is_test = true",
          f"{len(users)}/{len(USERS)}")


# ═══ Small readers ══════════════════════════════════════════════════════════
async def sig(conn, org, version_id) -> list[tuple]:
    rows = await scoped(conn, org, "SELECT global_security_id, target_weight, drift_min, drift_max, asset_class, "
                                   "resolution_state FROM portfolio.model_version_positions WHERE version_id = $1::uuid "
                                   "AND valid_to IS NULL ORDER BY row_num", str(version_id))
    return [(str(r["global_security_id"]) if r["global_security_id"] else None, D(r["target_weight"]),
             D(r["drift_min"]), D(r["drift_max"]), r["asset_class"], r["resolution_state"]) for r in rows]


async def version_row(conn, org, vid):
    return await scoped(conn, org, "SELECT * FROM portfolio.model_versions WHERE id = $1::uuid AND system_to IS NULL",
                        str(vid), fetch="fetchrow")


async def upload_outcome(conn, org, upload_id) -> str | None:
    row = await scoped(conn, org, "SELECT COALESCE((SELECT e.event FROM portfolio.model_upload_events e WHERE "
                                  "e.upload_id = u.id AND e.event IN ('draft_created','rejected','abandoned','failed') "
                                  "ORDER BY e.created_at DESC, e.id DESC LIMIT 1), u.outcome) AS o FROM "
                                  "portfolio.model_uploads u WHERE u.id = $1::uuid", str(upload_id), fetch="fetchrow")
    return row["o"] if row else None


async def confirm(uid, upload_id, **body):
    return await api(uid, "POST", f"/api/v1/models/uploads/{upload_id}/confirm", body)


async def to_draft(uid, name, data, model_name, **body):
    code, out = await upload(uid, name, data, provider=body.pop("provider", None))
    if code != 201 or not out.get("upload") or out.get("duplicate_of"):
        return code, out, None
    c2, o2 = await confirm(uid, out["upload"]["id"], model_name=model_name, **body)
    return c2, o2, out["upload"]["id"]


async def approve_path(model_id, version_id, maker=None, checker=None) -> tuple[int, dict]:
    maker, checker = maker or U_ADMIN1, checker or U_ADMIN2
    c1, o1 = await api(maker, "POST", f"/api/v1/models/{model_id}/versions/{version_id}/submit", {})
    if c1 != 200:
        return c1, o1
    return await api(checker, "POST", f"/api/v1/models/{model_id}/versions/{version_id}/approve", {})


async def task_for(conn, org, version_id) -> list:
    return await scoped(conn, org, "SELECT rs.id, rs.status, r.started_by, r.context->'human_task'->>'kind' AS kind, "
                                   "r.context->'human_task'->>'maker_id' AS maker_id, ws.assigned_permission "
                                   "FROM workflow_run_steps rs JOIN workflow_runs r ON r.id = rs.workflow_run_id "
                                   "JOIN workflow_steps ws ON ws.id = rs.workflow_step_id WHERE "
                                   "r.context->'human_task'->>'subject_id' = $1 ORDER BY r.started_at", str(version_id))


# ═══ Sections ═══════════════════════════════════════════════════════════════
async def y_task1(conn):
    section("[Y] Task 1 findings 1a–1e — files, signatures and the chosen patterns")
    import inspect

    from services import human_tasks as HT
    from services import models as M  # noqa: F401 — registers the model task kinds
    from services import replica_terms as T

    find("1a human tasks", "services/human_tasks.py: open_task(pool, *, org_id, kind, subject_type, subject_id, "
         "maker_id, eligible_permission, title, payload, action_path, actor_kind) / complete_task(pool, *, org_id, "
         "task_id, actor_id, outcome, note, actor_kind) / cancel_task(pool, *, org_id, task_id, actor_id, reason) / "
         "list_tasks(conn, org_id, *, status, kind, subject_type, subject_id, task_id) / get_task; subject = "
         "workflow_runs.context.human_task {kind, subject_type, subject_id, maker_id, title, payload, action_path}; "
         "outcome in workflow_run_steps.result {outcome, note, completed_by}. A module reacts through its kind's "
         "on_complete callback (idempotent) + a sync_*_task_outcome recovery. CHOSEN for rep3 and the rep1 move: "
         "on_complete — the only path that reacts however the task is completed; additive kwarg system_open lets "
         "the platform OPEN (never complete) a review task with no maker")
    params = inspect.signature(HT.register_kind).parameters
    kinds = HT.kinds()
    check("system_open" in params and kinds.get("model_version_approval") is not None
          and kinds["model_version_approval"].maker_checker is True
          and kinds.get("model_version_review") is not None and kinds["model_version_review"].system_open is True
          and kinds["model_version_review"].maker_checker is False
          and kinds.get("replica_terms_approval") is not None and kinds["replica_terms_approval"].maker_checker is False
          and all(kinds[k].on_complete is not None for k in ("model_version_approval", "model_version_review",
                                                            "replica_terms_approval")),
          "task kinds registered: model approval (maker-checker), model review (system-opened), terms approval")
    find("1b rep1 approval path", "services/replica_terms.py: _open_approval_todos (called by build_draft / "
         "correct_field / change_field) wrote member_todos for approve_replication holders; decide() closed them. "
         "verify_rep1.py asserted that path at its 'approval needed' check (~l.992) and 'approval tasks are closed' "
         "(~l.1026). MOVED: ensure_approval_task → HT.open_task (kind replica_terms_approval); decide → "
         "HT.complete_task, version change in on_complete; replica_terms_actions stays the domain log")
    src = (API_DIR / "services" / "replica_terms.py").read_text(encoding="utf-8")
    old = git_show(BEFORE_REF, "apps/api/services/replica_terms.py")
    rule = re.search(r'MAKER_CHECKER_RULE = \(([^)]*)\)', old or "")
    check("_open_approval_todos" not in src and "HT.complete_task" in src and "HT.open_task" in src
          and rule is not None and rule.group(0) in src,
          "rep1 approvals run through human tasks; the maker-checker refusal text is byte-identical to main")
    try:
        import xlrd  # noqa: F401
        xl = getattr(xlrd, "__version__", "?")
    except ImportError:
        xl = None
    import openpyxl as ox
    reqs = (API_DIR / "requirements.txt").read_text(encoding="utf-8")
    find("1c spreadsheet readers", f"openpyxl {ox.__version__} (installed, pinned 3.1.5); xlrd {xl or 'NOT installed'}"
         " (pinned xlrd==2.0.1 — pure Python, cached values only); pandas present but unused (not needed)")
    check("openpyxl==3.1.5" in reqs and "xlrd==2.0.1" in reqs, "openpyxl and xlrd pinned in requirements.txt")
    check(xl is not None, "xlrd is installed (pip install -r requirements.txt) — needed for the .xls proofs")
    pol = await platform(conn, "SELECT cmd, qual FROM pg_policies WHERE tablename = 'fee_schedules'")
    find("1d fee schedules", "referenced by fee_schedules.id of an APPROVED row (that row is never mutated; an edit "
         "writes version N+1 as a new DRAFT row with a new id) — models.default_fee_schedule_id FK, validated APPROVED "
         "+ same org in services.models.update_model; fee_schedules has an org-isolation ALL policy, so model code "
         "reads it on the RLS pool")
    check(len(pol) >= 1 and all("current_org_id" in (r["qual"] or "") for r in pol), "fee_schedules readable under RLS")
    find("1e platform rows", "rep2 custodian_capabilities: owner_scope 'platform', org_id NULL, SELECT USING true, "
         "writes only under app.is_super_admin. Library entries follow it for shelf 'platform'; the entry's content "
         "is a model owned by the Hollisworks org, admitted to every org read-only by a SECURITY DEFINER carve-out")


async def y_schema(conn):
    section("[Y] Schema landed: tables, constraints, partial uniques, triggers, RLS with a policy for every operation")
    code, out, err = run_cmd([sys.executable, str(HERE / "_apply_rep3.py"), "--verify", "--fixture-orgs",
                              ",".join(str(o) for o in FIXTURE_ORGS)], REPO, timeout=600)
    gaps = [ln for ln in out.splitlines() if ln.startswith("[GAP]")]
    check(code == 0 and "SCHEMA OK" in out and not gaps, "_apply_rep3.py --verify: SCHEMA OK (catalog read back)",
          f"{gaps[:3]} {err[-200:] if code else ''}")
    try:
        await scoped(conn, ORG_A, "INSERT INTO portfolio.models (org_id, name, source) VALUES ($1, 'Holdings of ZZRM', "
                                  "'upload')", ORG_A, fetch="execute")
        check(False, "the database refuses a 'Holdings of' name for an upload model")
    except Exception as exc:  # noqa: BLE001
        check("models_reserved_name_chk" in str(exc), "the database refuses a 'Holdings of' name for an upload model",
              str(exc)[:100])
    await scoped(conn, ORG_A, "INSERT INTO portfolio.models (org_id, name, source) VALUES ($1, 'Holdings of ZZRMF', "
                              "'fund_replica')", ORG_A, fetch="execute")
    n = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.models WHERE name = 'Holdings of ZZRMF'",
                     fetch="fetchval")
    check(int(n) == 1, "...and allows it for a fund_replica model (rep4's source)")


async def y_upload(conn):
    section("[Y] Upload: xlsx = xls = csv; untrusted-file rules; row-numbered messages; unresolved queued; duplicates")
    csv1 = csv_bytes(HEADER, ROWS_PCT, bom=True)
    code, out, up_csv = await to_draft(U_ADMIN1, "model.csv", csv1, "REP3 From CSV")
    check(code == 201 and out.get("version", {}).get("status") == "draft", "a CSV with a BOM parses into a draft",
          f"{code} {str(out)[:200]}")
    if code != 201:
        return
    S.update(m_csv=out["model"]["id"], v_csv=out["version"]["id"], up_csv=up_csv)
    basis = await scoped(conn, ORG_A, "SELECT weight_basis FROM portfolio.model_upload_events WHERE upload_id = $1::uuid "
                                      "AND event = 'draft_created'", up_csv, fetch="fetchval")
    check(basis == "percent", "percent weights detected and stated (weight_basis 'percent')", f"{basis}")
    sig_csv = await sig(conn, ORG_A, S["v_csv"])
    expect = [(S["sec_A"], Decimal("0.5"), Decimal("0.4"), Decimal("0.6"), "Equity", "resolved"),
              (S["sec_B"], Decimal("0.3"), None, None, "Equity", "resolved"),
              (S["sec_C"], Decimal("0.2"), None, None, "Bond", "resolved")]
    check(sig_csv == expect, "CSV draft: weights converted to fractions, bands kept, CUSIP and ticker rows resolved",
          f"{sig_csv}")
    xb = xlsx_bytes(HEADER, ROWS_PCT)
    code, out = await upload(U_ADMIN1, "model.xlsx", xb)
    sheets = [s["name"] for s in out.get("sheets", [])] if code == 201 else []
    check(code == 201 and sheets == ["Cover", "Holdings"], "a multi-sheet workbook lists its sheets", f"{code} {sheets}")
    if code == 201:
        up_x = out["upload"]["id"]
        c, pv = await api(U_ADMIN1, "POST", f"/api/v1/models/uploads/{up_x}/preview", {"sheet": "Cover"})
        check(c == 200 and any(m["code"] == "no_header" for m in pv["report"]["messages"]),
              "the cover sheet has no header row — stated, not guessed", f"{c}")
        c, cf = await confirm(U_ADMIN1, up_x, sheet="Holdings", model_name="REP3 From XLSX")
        check(c == 201, "the chosen sheet becomes a draft", f"{c} {str(cf)[:200]}")
        if c == 201:
            S.update(m_xlsx=cf["model"]["id"], v_xlsx=cf["version"]["id"])
            check(await sig(conn, ORG_A, S["v_xlsx"]) == sig_csv, ".xlsx draft is IDENTICAL to the .csv draft")
    try:
        import xlrd  # noqa: F401
        have_xlrd = True
    except ImportError:
        have_xlrd = False
    if have_xlrd:
        code, out, _ = await to_draft(U_ADMIN1, "model.xls", xls_bytes(HEADER, ROWS_PCT), "REP3 From XLS")
        check(code == 201, "a .xls (BIFF8) workbook parses into a draft", f"{code} {str(out)[:200]}")
        if code == 201:
            check(await sig(conn, ORG_A, out["version"]["id"]) == sig_csv, ".xls draft is IDENTICAL to the .csv draft")
    else:
        check(False, ".xls proof needs xlrd — run pip install -r apps/api/requirements.txt")
    frac = csv_bytes(HEADER, [[r[0], r[1], D(r[2]) / 100, D(r[3]) / 100 if r[3] else None,
                               D(r[4]) / 100 if r[4] else None, r[5]] for r in ROWS_PCT], bom=False, delimiter=";")
    code, out = await upload(U_ADMIN1, "fraction.csv", frac)
    if code == 201:
        c, pv = await api(U_ADMIN1, "POST", f"/api/v1/models/uploads/{out['upload']['id']}/preview", {})
        ws = [D(p["target_weight"]) for p in pv.get("report", {}).get("positions", [])]
        check(c == 200 and pv["report"]["weight_basis"] == "fraction" and ws == [Decimal("0.5"), Decimal("0.3"),
                                                                                 Decimal("0.2")],
              "fractions (sums ~1, ';'-delimited) detected as fractions — the same weights", f"{ws}")
    else:
        check(False, "fraction file uploaded", f"{code}")
    # a formula with no cached value
    wb = openpyxl.Workbook()
    ws_ = wb.active
    ws_.append(HEADER)
    ws_.append(ROWS_PCT[0])
    ws_.append([TICKER["B"], CUSIP["B"], "=100-50-20", None, None, "Equity"])
    ws_.append(ROWS_PCT[2])
    buf = io.BytesIO()
    wb.save(buf)
    code, out = await upload(U_ADMIN1, "formula.xlsx", buf.getvalue())
    if code == 201:
        up_f = out["upload"]["id"]
        c, cf = await confirm(U_ADMIN1, up_f, model_name="REP3 Formula")
        msgs = (cf.get("detail") or {}).get("messages", []) if isinstance(cf.get("detail"), dict) else []
        hit = [m for m in msgs if m["code"] == "formula_no_value"]
        check(c == 422 and len(hit) == 1 and hit[0]["row"] == 3 and hit[0]["severity"] == "error"
              and await upload_outcome(conn, ORG_A, up_f) == "rejected",
              "a formula cell with no cached value is a row-3 error (never evaluated) and the upload is logged rejected",
              f"{c} {hit}")
        S["up_rejected"] = up_f
    else:
        check(False, "formula workbook uploaded", f"{code}")
    puts0 = len(PUTS)
    code, out = await upload(U_ADMIN1, "macro.xlsm", xb)
    det = out.get("detail") or {}
    up_m = (det.get("upload") or {}).get("id") if isinstance(det, dict) else None
    row = await scoped(conn, ORG_A, "SELECT outcome, r2_key FROM portfolio.model_uploads WHERE id = $1::uuid",
                       up_m, fetch="fetchrow") if up_m else None
    check(code == 422 and det.get("messages", [{}])[0].get("code") == "macro_refused" and row is not None
          and row["outcome"] == "failed" and row["r2_key"] is None and len(PUTS) == puts0,
          ".xlsm refused (422), logged 'failed', never stored", f"{code} {dict(row) if row else None}")
    S["up_failed"] = up_m
    big = csv1 + b"#" * (5 * 1024 * 1024)
    code, out = await upload(U_ADMIN1, "big.csv", big)
    det = out.get("detail") or {}
    up_b = (det.get("upload") or {}).get("id") if isinstance(det, dict) else None
    size = await scoped(conn, ORG_A, "SELECT size_bytes FROM portfolio.model_uploads WHERE id = $1::uuid", up_b,
                        fetch="fetchval") if up_b else None
    check(code == 422 and det.get("messages", [{}])[0].get("code") == "too_large" and size == len(big),
          "an oversize file (> 5 MB) is refused and logged with its TRUE size", f"{code} {size} vs {len(big)}")
    many = csv_bytes(["Ticker", "Weight"], [[f"ZZRQ{i:05d}", "0.01"] for i in range(5100)], bom=False)
    code, out = await upload(U_ADMIN1, "many.csv", many)
    det = out.get("detail") or {}
    check(code == 422 and isinstance(det, dict) and det.get("messages", [{}])[0].get("code") == "too_many_rows",
          "more than 5,000 rows is refused (logged failed)", f"{code} {str(det)[:150]}")
    short = csv_bytes(HEADER, [ROWS_PCT[0], ROWS_PCT[1], [TICKER["C"], None, 19, None, None, "Bond"]])
    code, out = await upload(U_ADMIN1, "short.csv", short)
    if code == 201:
        up_s = out["upload"]["id"]
        c1, o1 = await confirm(U_ADMIN1, up_s, model_name="REP3 Short")
        m1 = (o1.get("detail") or {}).get("messages", []) if isinstance(o1.get("detail"), dict) else []
        c2, o2 = await confirm(U_ADMIN1, up_s, model_name="REP3 Short", fill_with_cash=True)
        cash = await scoped(conn, ORG_A, "SELECT target_weight FROM portfolio.model_version_positions WHERE version_id = "
                                         "$1::uuid AND resolution_state = 'cash' AND valid_to IS NULL",
                            o2.get("version", {}).get("id"), fetch="fetchval") if c2 == 201 else None
        check(c1 == 422 and any(m["code"] == "weight_sum" for m in m1) and c2 == 201 and D(cash) == Decimal("0.01"),
              "weights summing to 99% are refused — unless 'fill with cash' (then a 1% cash target)", f"{c1} {c2} {cash}")
    else:
        check(False, "short file uploaded", f"{code}")
    bad = csv_bytes(HEADER, [[TICKER["A"], CUSIP["A"], 60, None, None, None], [TICKER["A"], CUSIP["A"], 50, None, None,
                                                                                None],
                             [TICKER["B"], CUSIP["B"], -10, None, None, None]])
    code, out = await upload(U_ADMIN1, "bad.csv", bad)
    if code == 201:
        c, o = await confirm(U_ADMIN1, out["upload"]["id"], model_name="REP3 Bad")
        ms = (o.get("detail") or {}).get("messages", []) if isinstance(o.get("detail"), dict) else []
        neg = [m for m in ms if m["code"] == "weight_negative"]
        dup = [m for m in ms if m["code"] == "duplicate_security"]
        check(c == 422 and len(neg) == 1 and neg[0]["row"] == 4 and len(dup) == 1 and dup[0]["row"] == 3,
              "a negative weight (row 4) and a duplicate security (row 3) are refused, each with its row", f"{c} {ms[:4]}")
    else:
        check(False, "bad file uploaded", f"{code}")
    unres = csv_bytes(HEADER, [[TICKER["D"], MISSING_CUSIP, 40, None, None, None],
                               [TICKER["A"], CUSIP["A"], 60, None, None, None]])
    code, out, up_u = await to_draft(U_ADMIN1, "unresolved.csv", unres, "REP3 Unresolved")
    if code == 201:
        S.update(m_unres=out["model"]["id"], v_unres=out["version"]["id"], up_unres=up_u)
        s_ = await sig(conn, ORG_A, S["v_unres"])
        ids = {x[0] for x in s_}
        msgs = out.get("report", {}).get("messages", [])
        check(s_[0][0] is None and s_[0][5] == "unresolved" and S["sec_D"] not in ids
              and any(m["code"] == "unresolved" and m["row"] == 2 for m in msgs),
              "an unresolvable CUSIP is QUEUED unresolved — never matched to the decoy holding its ticker", f"{s_}")
        trad = (await version_row(conn, ORG_A, S["v_unres"]))["tradability"]
        check("fractional_shares unknown at altruist" in str(trad),
              "tradability noted per custodian the org uses (informational): fractional_shares unknown at altruist",
              str(trad)[:160])
    else:
        check(False, "unresolved file becomes a draft (unresolved rows can be saved)", f"{code} {str(out)[:200]}")
    puts1 = len(PUTS)
    code, out = await upload(U_ADMIN1, "model-again.csv", csv1)
    dup_id = out.get("upload", {}).get("id") if code == 201 else None
    row = await scoped(conn, ORG_A, "SELECT outcome, duplicate_of, r2_key FROM portfolio.model_uploads WHERE id = "
                                    "$1::uuid", dup_id, fetch="fetchrow") if dup_id else None
    orig = await scoped(conn, ORG_A, "SELECT r2_key FROM portfolio.model_uploads WHERE id = $1::uuid", up_csv,
                        fetch="fetchval")
    check(code == 201 and out.get("duplicate_of", {}).get("id") == up_csv and row is not None
          and row["outcome"] == "duplicate" and str(row["duplicate_of"]) == up_csv and row["r2_key"] == orig
          and len(PUTS) == puts1 and len([k for k in STORE if k.endswith(orig.rsplit("/", 1)[-1])]) == 1,
          "the same file again → a new log row 'duplicate' (duplicate_of the first), ONE R2 object",
          f"{code} {dict(row) if row else None}")


async def y_mapping(conn):
    section("[Y] Mapping suggestions and saved provider mappings")
    code, out = await upload(U_ADMIN1, "synonyms.csv", csv_bytes(["Symbol", "Allocation"], [[TICKER["A"], 100]]))
    sugg = {s["header"]: (s["field"], s["source"]) for s in out.get("suggestions", [])} if code == 201 else {}
    check(sugg.get("Symbol") == ("ticker", "synonym") and sugg.get("Allocation") == ("weight", "synonym"),
          "synonyms auto-suggested: Symbol → Ticker, Allocation → Target Weight", f"{sugg}")
    prov = "REP3 Verify Provider"
    f1 = csv_bytes(["Fund Code", "Model Pct"], [[TICKER["A"], 60], [TICKER["B"], 40]])
    code, out = await upload(U_ADMIN1, "prov1.csv", f1, provider=prov)
    first = {s["header"]: s["field"] for s in out.get("suggestions", [])} if code == 201 else {"x": "x"}
    check(code == 201 and first.get("Fund Code") is None and first.get("Model Pct") is None,
          "unknown headers get no suggestion before a mapping is saved", f"{first}")
    if code != 201:
        return
    c, m = await api(U_ADMIN1, "POST", "/api/v1/models/mappings", {"provider_name": prov, "mapping": {
        "Fund Code": "ticker", "Model Pct": "weight"}, "upload_id": out["upload"]["id"]})
    row = await scoped(conn, ORG_A, "SELECT id, mapping, version_no FROM portfolio.model_column_mappings WHERE "
                                    "provider_key = lower($1) AND valid_to IS NULL", prov, fetch="fetchrow")
    logged = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_upload_events WHERE upload_id = $1::uuid "
                                       "AND event = 'mapping_saved'", out["upload"]["id"], fetch="fetchval")
    evt = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_events WHERE action = 'save_mapping'",
                       fetch="fetchval")
    check(c == 200 and row is not None and json.loads(row["mapping"]) == {"Fund Code": "ticker", "Model Pct": "weight"}
          and row["version_no"] == 1 and int(logged) == 1 and int(evt) >= 1,
          "saving the mapping is a logged action (upload event + model event); re-read version 1", f"{c}")
    f2 = csv_bytes(["Fund Code", "Model Pct"], [[TICKER["A"], 55], [TICKER["B"], 45]])
    code, out2 = await upload(U_ADMIN1, "prov2.csv", f2, provider=prov)
    second = {s["header"]: (s["field"], s["source"]) for s in out2.get("suggestions", [])} if code == 201 else {}
    check(second.get("Fund Code") == ("ticker", "saved") and second.get("Model Pct") == ("weight", "saved"),
          "the provider's next upload applies the saved mapping first", f"{second}")
    if code == 201:
        c, cf = await confirm(U_ADMIN1, out2["upload"]["id"], provider_name=prov, model_name="REP3 Provider Model")
        mid = await scoped(conn, ORG_A, "SELECT mapping_id FROM portfolio.model_upload_events WHERE upload_id = $1::uuid "
                                        "AND event = 'draft_created'", out2["upload"]["id"], fetch="fetchval")
        check(c == 201 and row is not None and mid == row["id"],
              "...and the draft records which saved mapping was used", f"{c} {mid}")


async def y_log(conn):
    section("[Y] Upload log complete for every outcome, immutable facts, safe CSV export")
    code, out = await upload(U_ADMIN1, "abandon-me.csv", csv_bytes(HEADER, [[TICKER["A"], CUSIP["A"], 100, None, None,
                                                                             "Equity"]]))
    if code == 201:
        c, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/uploads/{out['upload']['id']}/abandon",
                         {"reason": "verify: not needed"})
        check(c == 200 and await upload_outcome(conn, ORG_A, out["upload"]["id"]) == "abandoned",
              "an upload abandoned before a draft is logged 'abandoned' with its reason", f"{c}")
    code, out = await upload(U_ADMIN1, "=HYPERLINK(1).csv", csv_bytes(["Ticker", "Weight"], [[TICKER["A"], 100]]))
    c, lst = await api(U_ADMIN1, "GET", "/api/v1/models/uploads")
    rows = lst.get("rows", [])
    outs = {r["outcome"] for r in rows}
    check(c == 200 and {"draft_created", "rejected", "abandoned", "duplicate", "failed"} <= outs,
          "the log holds every outcome: draft_created, rejected, abandoned, duplicate, failed", f"{sorted(outs)}")
    rej = next((r for r in rows if r["id"] == S.get("up_rejected")), None)
    check(rej is not None and bool(rej["messages"]) and all(isinstance(m.get("row"), int) for m in rej["messages"]
                                                           if m["code"] == "formula_no_value"),
          "a rejected upload's messages carry row numbers", f"{(rej or {}).get('messages')}")
    before = await scoped(conn, ORG_A, "SELECT file_name, sha256, outcome, size_bytes FROM portfolio.model_uploads WHERE "
                                       "id = $1::uuid", S["up_csv"], fetch="fetchrow")
    for sql, label in (("UPDATE portfolio.model_uploads SET file_name = 'x' WHERE id = $1::uuid", "model_uploads"),
                       ("UPDATE portfolio.model_upload_events SET reason = 'x' WHERE upload_id = $1::uuid",
                        "model_upload_events")):
        try:
            await scoped(conn, ORG_A, sql, S["up_csv"], fetch="execute")
            check(False, f"UPDATE of {label} refused")
        except Exception as exc:  # noqa: BLE001
            check("permanent record" in str(exc), f"UPDATE of {label} refused (facts never change)", str(exc)[:80])
    after = await scoped(conn, ORG_A, "SELECT file_name, sha256, outcome, size_bytes FROM portfolio.model_uploads WHERE "
                                      "id = $1::uuid", S["up_csv"], fetch="fetchrow")
    check(before is not None and dict(before) == dict(after), "the upload's facts are unchanged (re-read)")
    c, content, headers = await api_raw(U_ADMIN1, "GET", "/api/v1/models/uploads/export.csv")
    text = content.decode("utf-8") if c == 200 else ""
    parsed = list(csv.reader(io.StringIO(text)))
    total = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_uploads", fetch="fetchval")
    cells = [x for r in parsed[1:] for x in r]
    check(c == 200 and "text/csv" in headers.get("content-type", "") and len(parsed) - 1 == int(total)
          and "'=HYPERLINK(1).csv" in cells and not any(x[:1] in ("=", "+", "-", "@", "\t", "\r") for x in cells),
          "CSV export: every row, the =cmd file name escaped with an apostrophe, no cell begins with = + - @ tab CR",
          f"{c} {len(parsed) - 1}/{total}")


async def y_versions(conn):
    section("[Y] Versions: exact diff; maker-checker on the identical request; unresolved blocks approval; supersede; "
            "event; immutability")
    from services import human_tasks as HT
    from services.database import get_pool

    m, v1 = S.get("m_csv"), S.get("v_csv")
    if not (m and v1):
        check(False, "a CSV model exists for the version proofs")
        return
    c, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/{m}/versions/{v1}/submit", {})
    t = await task_for(conn, ORG_A, v1)
    check(c == 200 and len(t) == 1 and t[0]["kind"] == "model_version_approval" and t[0]["started_by"] == U_ADMIN1
          and t[0]["assigned_permission"] == "approve_models" and t[0]["status"] == "active",
          "submit opens ONE human task (approve_models) whose engine maker is the submitter", f"{c} {[dict(x) for x in t]}")
    path = f"/api/v1/models/{m}/versions/{v1}/approve"
    c1, o1 = await api(U_ADMIN1, "POST", path, {"note": "ok"})
    st = (await version_row(conn, ORG_A, v1))["status"]
    check(c1 == 403 and "Maker-checker" in str(o1.get("detail")) and st == "submitted",
          "the MAKER's approve is refused (403 naming the rule); the version is unchanged", f"{c1} {o1}")
    try:
        await scoped(conn, ORG_A, "UPDATE portfolio.model_versions SET status = 'approved', checker_id = $2, approved_at "
                                  "= now() WHERE id = $1::uuid", v1, U_ADMIN1, fetch="execute")
        check(False, "the database refuses a maker as checker")
    except Exception as exc:  # noqa: BLE001
        check("maker_checker" in str(exc), "the database CHECK refuses a maker as checker", str(exc)[:90])
    c2, o2 = await api(U_ADMIN2, "POST", path, {"note": "ok"})
    row = await version_row(conn, ORG_A, v1)
    check(c2 == 200 and row["status"] == "approved" and row["checker_id"] == U_ADMIN2,
          "the identical request by a different approver succeeds", f"{c2} {str(o2)[:160]}")
    ev = await scoped(conn, ORG_A, "SELECT payload FROM domain_events WHERE event_type = 'model.version_approved' AND "
                                   "source_id = $1::uuid", v1)
    check(len(ev) == 1 and "diff" in json.loads(ev[0]["payload"]), "'model.version_approved' emitted with the diff")
    pool = await get_pool()
    with _ctx(ORG_A):
        try:
            await HT.open_task(pool, org_id=ORG_A, kind="model_version_approval", subject_type="model_version",
                               subject_id=v1, maker_id=None, eligible_permission="approve_models", title="x",
                               actor_kind="system")
            check(False, "the platform cannot open a maker-checker task")
        except HT.AgentActorRefused:
            check(True, "the platform cannot open a maker-checker task (system opens review kinds only)")
        try:
            await HT.open_task(pool, org_id=ORG_A, kind="model_version_review", subject_type="model_version",
                               subject_id=v1, maker_id=U_ADMIN1, eligible_permission="approve_models", title="x",
                               actor_kind="agent")
            check(False, "an agent cannot open a task")
        except HT.AgentActorRefused:
            check(True, "an agent cannot open a task")
    c, d = await api(U_ADMIN1, "POST", f"/api/v1/models/{m}/drafts", {})
    v2 = d.get("version", {}).get("id") if c == 201 else None
    if not v2:
        check(False, "a new draft from the approved version", f"{c} {d}")
        return
    S["v_csv2"] = v2
    c, e = await api(U_ADMIN1, "POST", f"/api/v1/models/{m}/versions/{v2}/positions", {"changes": [
        {"op": "upsert", "row_num": 2, "weight_pct": "45", "drift_min_pct": "40", "drift_max_pct": "50"},
        {"op": "upsert", "row_num": 3, "weight_pct": "35"},
        {"op": "remove", "row_num": 4},
        {"op": "upsert", "raw_cusip": CUSIP["F"], "weight_pct": "20"}]})
    diff = e.get("version", {}).get("diff", {}) if c == 200 else {}
    ch = {x["key"]: D(x["delta"]) for x in diff.get("changed", [])}
    check(c == 200 and ch == {S["sec_A"]: Decimal("-0.05"), S["sec_B"]: Decimal("0.05")}
          and [x["key"] for x in diff.get("added", [])] == [S["sec_F"]]
          and [x["key"] for x in diff.get("removed", [])] == [S["sec_C"]]
          and D(diff.get("turnover")) == Decimal("0.25"),
          "diff exact: added F, removed C, A −5%, B +5%, turnover 25%", f"{c} {diff}")
    c, o = await approve_path(m, v2)
    vers = await scoped(conn, ORG_A, "SELECT id, status FROM portfolio.model_versions WHERE model_id = $1::uuid AND "
                                     "system_to IS NULL", m)
    stt = {str(r["id"]): r["status"] for r in vers}
    check(c == 200 and stt.get(v1) == "superseded" and stt.get(v2) == "approved"
          and list(stt.values()).count("approved") == 1,
          "approval supersedes the prior version; exactly ONE approved version live (re-read)", f"{stt}")
    ev2 = await scoped(conn, ORG_A, "SELECT payload FROM domain_events WHERE event_type = 'model.version_approved' AND "
                                    "source_id = $1::uuid", v2, fetch="fetchval")
    check(ev2 is not None and D(json.loads(ev2)["diff"]["turnover"]) == Decimal("0.25"),
          "...and its event carries the exact diff (turnover 25%)")
    snap = await sig(conn, ORG_A, v2)
    c, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/{m}/versions/{v2}/positions",
                     {"changes": [{"op": "upsert", "row_num": 2, "weight_pct": "1"}]})
    check(c == 409, "editing an APPROVED version is 409", f"{c}")
    try:
        await scoped(conn, ORG_A, "UPDATE portfolio.model_versions SET notes = 'x' WHERE id = $1::uuid", v2,
                     fetch="execute")
        check(False, "a direct UPDATE of an approved version is refused")
    except Exception as exc:  # noqa: BLE001
        check("never changed" in str(exc), "a direct UPDATE of an approved version is refused (trigger)", str(exc)[:80])
    try:
        await scoped(conn, ORG_A, "INSERT INTO portfolio.model_version_positions (org_id, version_id, row_num, "
                                  "target_weight, resolution_state) VALUES ($1, $2::uuid, 99, 0, 'cash')", ORG_A, v2,
                     fetch="execute")
        check(False, "a position cannot be added to an approved version")
    except Exception as exc:  # noqa: BLE001
        check("only while it is a draft" in str(exc), "a position cannot be added to an approved version (trigger)")
    check(await sig(conn, ORG_A, v2) == snap, "the approved version's positions are unchanged (re-read)")
    mu, vu = S.get("m_unres"), S.get("v_unres")
    if mu and vu:
        c, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/{mu}/versions/{vu}/submit", {})
        c2, o2 = await api(U_ADMIN2, "POST", f"/api/v1/models/{mu}/versions/{vu}/approve", {})
        check(c == 200 and c2 == 409 and "unresolved" in str(o2.get("detail"))
              and (await version_row(conn, ORG_A, vu))["status"] == "submitted",
              "a version with unresolved rows cannot be approved (409 naming them)", f"{c} {c2}")
        c3, _ = await api(U_ADMIN2, "POST", f"/api/v1/models/{mu}/versions/{vu}/reject", {})
        c4, _ = await api(U_ADMIN2, "POST", f"/api/v1/models/{mu}/versions/{vu}/reject", {"note": "resolve row 2"})
        r = await version_row(conn, ORG_A, vu)
        check(c3 == 422 and c4 == 200 and r["status"] == "draft" and r["decision_note"] == "resolve row 2",
              "reject needs a note (422 without); with one the version returns to draft", f"{c3} {c4}")
        c5, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/{mu}/versions/{vu}/abandon", {"reason": "verify"})
        logged = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_events WHERE version_id = $1::uuid AND "
                                           "action = 'abandon'", vu, fetch="fetchval")
        check(c5 == 200 and (await version_row(conn, ORG_A, vu))["status"] == "abandoned" and int(logged) == 1
              and await upload_outcome(conn, ORG_A, S["up_unres"]) == "abandoned",
              "abandon is logged and the upload's outcome becomes abandoned", f"{c5}")
    else:
        check(False, "the unresolved model exists for the approval-block proof")


async def y_names(conn):
    section("[Y] \"Holdings of\" reserved")
    code, out = await upload(U_ADMIN1, "reserved.csv", csv_bytes(["Ticker", "Weight"], [[TICKER["B"], 100]]))
    if code == 201:
        c, o = await confirm(U_ADMIN1, out["upload"]["id"], model_name="Holdings of SPY")
        check(c == 422 and "reserved" in str(o.get("detail")), "an upload model named 'Holdings of …' is refused (422)",
              f"{c}")
    c, o = await api(U_ADMIN1, "POST", "/api/v1/models", {"name": "holdings of things", "sleeves": [
        {"child_model_id": str(S.get("m_csv")), "weight_pct": "100"}]})
    check(c == 422 and "reserved" in str(o.get("detail")), "a composite named 'holdings of …' is refused (422)", f"{c}")


async def y_sleeves(conn):
    section("[Y] Sleeves: look-through exact, no cycles, follow-live drafts + review tasks")
    m1, m2 = S.get("m_csv"), S.get("m_xlsx")
    if not (m1 and m2):
        check(False, "two approved child models exist")
        return
    c, _ = await approve_path(m2, S["v_xlsx"])
    check(c == 200, "child model 2 approved", f"{c}")
    c, o = await api(U_ADMIN1, "POST", "/api/v1/models", {"name": "REP3 Composite Live", "sleeves": [
        {"child_model_id": m1, "weight_pct": "60", "follow": "live"},
        {"child_model_id": m2, "weight_pct": "40", "follow": "live"}]})
    if c != 201:
        check(False, "composite created", f"{c} {o}")
        return
    p1, pv1 = o["model"]["id"], o["version"]["id"]
    c, d = await api(U_ADMIN1, "GET", f"/api/v1/models/{p1}/versions/{pv1}")
    lt = {x["key"]: D(x["weight"]) for x in d.get("look_through") or []}
    expect = {S["sec_A"]: Decimal("0.47"), S["sec_B"]: Decimal("0.33"), S["sec_F"]: Decimal("0.12"),
              S["sec_C"]: Decimal("0.08")}
    check(c == 200 and lt == expect and sum(lt.values()) == Decimal(1),
          "look-through = sleeve weight × child weights, merged by security (exact)", f"{lt}")
    c, o = await api(U_ADMIN1, "POST", "/api/v1/models", {"name": "REP3 Composite Pinned", "sleeves": [
        {"child_model_id": m1, "weight_pct": "100", "follow": "pinned", "pinned_version_id": S["v_csv2"]}]})
    pp = o.get("model", {}).get("id") if c == 201 else None
    c1, _ = await approve_path(p1, pv1)
    c2, _ = await approve_path(pp, o["version"]["id"]) if pp else (0, {})
    check(c1 == 200 and c2 == 200, "the live-following and the pinned composite are approved", f"{c1} {c2}")
    c, o = await api(U_ADMIN1, "POST", "/api/v1/models", {"name": "REP3 Composite Outer", "sleeves": [
        {"child_model_id": p1, "weight_pct": "100", "follow": "live"}]})
    check(c == 201, "a composite of a composite is allowed", f"{c}")
    c, d = await api(U_ADMIN1, "POST", f"/api/v1/models/{p1}/drafts", {})
    pd = d.get("version", {}).get("id") if c == 201 else None
    outer = o.get("model", {}).get("id")
    if pd and outer:
        c3, o3 = await api(U_ADMIN1, "PUT", f"/api/v1/models/{p1}/versions/{pd}/sleeves", {"sleeves": [
            {"child_model_id": outer, "weight_pct": "100", "follow": "live"}]})
        c4, o4 = await api(U_ADMIN1, "PUT", f"/api/v1/models/{p1}/versions/{pd}/sleeves", {"sleeves": [
            {"child_model_id": p1, "weight_pct": "100", "follow": "live"}]})
        n = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_sleeves WHERE parent_version_id = $1::uuid "
                                      "AND valid_to IS NULL AND child_model_id = ANY($2::uuid[])", pd, [outer, p1],
                         fetch="fetchval")
        check(c3 == 409 and "cycle" in str(o3.get("detail")) and c4 == 409 and int(n) == 0,
              "a cycle is refused (outer contains this one; a model containing itself), nothing written",
              f"{c3} {c4}")
        await api(U_ADMIN1, "POST", f"/api/v1/models/{p1}/versions/{pd}/abandon", {"reason": "verify cycle"})
    else:
        check(False, "a draft of the composite for the cycle proof", f"{c}")
    pinned_before = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_versions WHERE model_id = $1::uuid "
                                              "AND system_to IS NULL", pp, fetch="fetchval") if pp else None
    c, d = await api(U_ADMIN1, "POST", f"/api/v1/models/{m1}/drafts", {})
    nv = d.get("version", {}).get("id") if c == 201 else None
    if not nv:
        check(False, "a new child version", f"{c}")
        return
    await api(U_ADMIN1, "POST", f"/api/v1/models/{m1}/versions/{nv}/positions", {"changes": [
        {"op": "upsert", "row_num": 2, "weight_pct": "40", "drift_min_pct": "35", "drift_max_pct": "45"},
        {"op": "upsert", "row_num": 3, "weight_pct": "40"}]})
    c, o = await approve_path(m1, nv)
    fu = o.get("version", {}).get("follow_up", {}) if c == 200 else {}
    drafted = [x for x in fu.get("composite_drafts", []) if x.get("model_id") == p1]
    pv = await scoped(conn, ORG_A, "SELECT id, status, origin, maker_ids FROM portfolio.model_versions WHERE model_id = "
                                   "$1::uuid AND status IN ('draft','submitted') AND system_to IS NULL", p1)
    tsk = await task_for(conn, ORG_A, pv[0]["id"]) if pv else []
    pinned_after = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_versions WHERE model_id = $1::uuid AND "
                                             "system_to IS NULL", pp, fetch="fetchval") if pp else None
    check(c == 200 and len(drafted) == 1 and len(pv) == 1 and pv[0]["origin"] == "composite_follow"
          and pv[0]["status"] == "submitted" and list(pv[0]["maker_ids"]) == [] and len(tsk) == 1
          and tsk[0]["kind"] == "model_version_review" and tsk[0]["started_by"] is None and tsk[0]["maker_id"] is None,
          "a child's new approved version drafts the live-following composite with a review task (no maker)",
          f"{c} {fu}")
    check(pinned_before is not None and pinned_before == pinned_after, "...and nothing for the pinned composite",
          f"{pinned_before} → {pinned_after}")
    if pv:
        c, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/{p1}/versions/{pv[0]['id']}/approve", {})
        check(c == 200, "a system-imported draft is approved by an org approver (maker = system import)", f"{c}")


async def y_library(conn):
    section("[Y] Library: shelves, visibility, copy vs subscribe, agreement gate")
    from services import models as M
    from services.database import get_pool

    m2 = S.get("m_xlsx")
    if not m2:
        check(False, "a model to publish")
        return
    c, o = await api(U_ADMIN1, "POST", "/api/v1/models/library", {"model_id": m2, "shelf": "org",
                     "title": "REP3 Org Entry", "provider": "REP3", "update_cadence": "monthly",
                     "requires_agreement": True})
    eo = o.get("entry", {}).get("id") if c == 201 else None
    check(c == 201, "an org admin publishes an ORG-shelf entry", f"{c} {o}")
    if not eo:
        return
    _, lb = await api(U_ADMIN_B1, "GET", "/api/v1/models/library")
    cb, _ = await api(U_ADMIN_B1, "POST", f"/api/v1/models/library/{eo}/copy", {"name": "B steals"})
    check(eo not in {r["id"] for r in lb.get("rows", [])} and cb == 404,
          "the org entry is invisible to the other org (list and copy 404)", f"{cb}")
    c1, o1 = await api(U_ADMIN1, "POST", f"/api/v1/models/library/{eo}/copy", {"name": "REP3 Copy"})
    c2, _ = await api(U_VIEWER, "POST", f"/api/v1/models/library/{eo}/agreement", {})
    check(c1 == 409 and "agreement" in str(o1.get("detail")) and c2 == 403,
          "requires_agreement blocks copy (409) until confirmed; a viewer cannot confirm (403)", f"{c1} {c2}")
    c3, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/library/{eo}/agreement", {"note": "on file"})
    logged = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_events WHERE library_entry_id = $1::uuid AND "
                                       "action = 'agreement_confirmed' AND actor_id = $2", eo, U_ADMIN1, fetch="fetchval")
    check(c3 == 200 and int(logged) == 1, "an org admin's 'agreement on file' is logged", f"{c3}")
    c4, cp = await api(U_ADMIN1, "POST", f"/api/v1/models/library/{eo}/copy", {"name": "REP3 Copy"})
    c5, sb = await api(U_ADMIN1, "POST", f"/api/v1/models/library/{eo}/subscribe", {"name": "REP3 Subscribed"})
    check(c4 == 201 and cp["model"]["source"] == "library_copy" and cp["version"]["status"] == "draft"
          and c5 == 201 and sb["model"]["source"] == "library_subscription" and sb["version"]["status"] == "submitted",
          "copy → an independent draft; subscribe → a draft awaiting review", f"{c4} {c5}")
    c, _ = await api(U_ADMIN2, "POST", f"/api/v1/models/{sb['model']['id']}/versions/{sb['version']['id']}/approve", {})
    copy_n = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_versions WHERE model_id = $1::uuid AND "
                                       "system_to IS NULL", cp["model"]["id"], fetch="fetchval")
    c, d = await api(U_ADMIN1, "POST", f"/api/v1/models/{m2}/drafts", {})
    nv = d.get("version", {}).get("id") if c == 201 else None
    if nv:
        await api(U_ADMIN1, "POST", f"/api/v1/models/{m2}/versions/{nv}/positions", {"changes": [
            {"op": "upsert", "row_num": 2, "weight_pct": "55", "drift_min_pct": "40", "drift_max_pct": "60"},
            {"op": "upsert", "row_num": 3, "weight_pct": "25"}]})
        c, o = await approve_path(m2, nv)
        fu = o.get("version", {}).get("follow_up", {}) if c == 200 else {}
        subd = [x for x in fu.get("subscription_drafts", []) if x.get("model_id") == sb["model"]["id"]]
        copy_after = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.model_versions WHERE model_id = $1::uuid "
                                               "AND system_to IS NULL", cp["model"]["id"], fetch="fetchval")
        check(c == 200 and len(subd) == 1 and int(copy_after) == int(copy_n),
              "a new entry version drafts the SUBSCRIBER (with a task); the copy is untouched", f"{fu}")
    # ── platform shelf. The platform model is SEEDED directly (no human task is opened in the real
    # Hollisworks org): draft → submitted → approved, the guard trigger's only path. The 'submitted' state's
    # NOT NULL approval_task_id borrows an already-decided fixture task of org A (FK only).
    seed_task = await scoped(conn, ORG_A, "SELECT approval_task_id FROM portfolio.model_versions WHERE id = $1::uuid",
                             S["v_csv"], fetch="fetchval")
    if seed_task is None:
        check(False, "a decided fixture task exists to seed the platform model")
        return
    plat_m, plat_v1 = uuid.uuid4(), uuid.uuid4()
    S["plat_m"] = str(plat_m)
    async with conn.transaction():
        await conn.execute("SELECT set_config('app.is_super_admin', 'true', true)")
        await conn.execute("INSERT INTO portfolio.models (id, org_id, name, source, status, created_by) VALUES "
                           "($1, $2, 'REP3 Verify Platform Model', 'upload', 'active', $3)", plat_m, PLATFORM_ORG, U_PLAT)
        await conn.execute("INSERT INTO portfolio.model_versions (id, org_id, model_id, version_no, status, origin, "
                           "maker_ids, checker_id, approved_at, created_by) VALUES ($1, $2, $3, 1, 'draft', 'manual', "
                           "$4::uuid[], NULL, NULL, $5)", plat_v1, PLATFORM_ORG, plat_m, [U_PLAT], U_PLAT)
        for i, (k, w) in enumerate((("A", "0.6"), ("B", "0.4")), start=2):
            await conn.execute("INSERT INTO portfolio.model_version_positions (org_id, version_id, row_num, "
                               "global_security_id, raw_cusip, target_weight, resolution_state) VALUES ($1, $2, $3, "
                               "$4::uuid, $5, $6::numeric, 'resolved')", PLATFORM_ORG, plat_v1, i, S[f"sec_{k}"],
                               CUSIP[k], w)
        await conn.execute("UPDATE portfolio.model_versions SET status = 'submitted', approval_task_id = $2 WHERE id = $1",
                           plat_v1, seed_task)
        await conn.execute("UPDATE portfolio.model_versions SET status = 'approved', checker_id = $2, approved_at = "
                           "now() WHERE id = $1", plat_v1, U_SUPER)
    body = {"model_id": str(plat_m), "shelf": "platform", "title": "REP3 Platform Entry", "provider": "Hollisworks",
            "update_cadence": "quarterly", "licensing_note": "verify only"}
    c0, _ = await api(U_ADMIN1, "POST", "/api/v1/models/library", body)
    c1, o1 = await api(U_PLAT, "POST", "/api/v1/models/library", body)
    pe = o1.get("entry", {}).get("id") if c1 == 201 else None
    check(c0 == 403 and c1 == 201, "only Super Admin publishes to the PLATFORM shelf (org admin 403)", f"{c0} {c1}")
    if not pe:
        return
    _, la = await api(U_ADMIN1, "GET", "/api/v1/models/library")
    _, lb = await api(U_ADMIN_B1, "GET", "/api/v1/models/library")
    check(pe in {r["id"] for r in la.get("rows", [])} and pe in {r["id"] for r in lb.get("rows", [])},
          "the platform entry is visible to BOTH orgs")
    c2, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/library/{pe}/status", {"status": "unpublished", "reason": "x"})
    st = await scoped(conn, ORG_A, "UPDATE portfolio.model_library_entries SET title = 'hijack' WHERE id = $1::uuid",
                      pe, fetch="execute")
    st_b = await scoped(conn, ORG_B, "UPDATE portfolio.model_library_entries SET title = 'hijack' WHERE id = $1::uuid",
                        pe, fetch="execute")
    title = await platform(conn, "SELECT title FROM portfolio.model_library_entries WHERE id = $1::uuid", pe,
                           fetch="fetchval")
    check(c2 == 403 and st == "UPDATE 0" and st_b == "UPDATE 0" and title == "REP3 Platform Entry",
          "...and editable by NEITHER org (API 403; SQL UPDATE matches 0 rows under RLS)", f"{c2} {st} {st_b}")
    c3, cpb = await api(U_ADMIN_B1, "POST", f"/api/v1/models/library/{pe}/copy", {"name": "B Platform Copy"})
    c4, sba = await api(U_ADMIN1, "POST", f"/api/v1/models/library/{pe}/subscribe", {"name": "A Platform Sub"})
    check(c3 == 201 and c4 == 201, "org B copies and org A subscribes to the platform entry", f"{c3} {c4}")
    if not (c3 == 201 and c4 == 201):
        return
    await api(U_ADMIN2, "POST", f"/api/v1/models/{sba['model']['id']}/versions/{sba['version']['id']}/approve", {})
    plat_v2 = uuid.uuid4()
    async with conn.transaction():
        await conn.execute("SELECT set_config('app.is_super_admin', 'true', true)")
        await conn.execute("UPDATE portfolio.model_versions SET status = 'superseded' WHERE id = $1", plat_v1)
        await conn.execute("INSERT INTO portfolio.model_versions (id, org_id, model_id, version_no, status, origin, "
                           "maker_ids, created_by, parent_version_id) VALUES ($1, $2, $3, 2, 'draft', 'manual', "
                           "$4::uuid[], $5, $6)", plat_v2, PLATFORM_ORG, plat_m, [U_PLAT], U_PLAT, plat_v1)
        for i, (k, w) in enumerate((("A", "0.55"), ("B", "0.45")), start=2):
            await conn.execute("INSERT INTO portfolio.model_version_positions (org_id, version_id, row_num, "
                               "global_security_id, raw_cusip, target_weight, resolution_state) VALUES ($1, $2, $3, "
                               "$4::uuid, $5, $6::numeric, 'resolved')", PLATFORM_ORG, plat_v2, i, S[f"sec_{k}"],
                               CUSIP[k], w)
        await conn.execute("UPDATE portfolio.model_versions SET status = 'submitted', approval_task_id = $2 WHERE id = $1",
                           plat_v2, seed_task)
        await conn.execute("UPDATE portfolio.model_versions SET status = 'approved', checker_id = $2, approved_at = "
                           "now() WHERE id = $1", plat_v2, U_SUPER)
    pool = await get_pool()
    with _ctx(PLATFORM_ORG, True):
        fu = await M.after_approval(pool, org_id=str(PLATFORM_ORG), model_id=str(plat_m), version_id=str(plat_v2))
    sa = [x for x in fu["subscription_drafts"] if x.get("model_id") == sba["model"]["id"]]
    nb = await scoped(conn, ORG_B, "SELECT count(*) FROM portfolio.model_versions WHERE model_id = $1::uuid AND "
                                   "system_to IS NULL", cpb["model"]["id"], fetch="fetchval")
    tk = await task_for(conn, ORG_A, sa[0]["version_id"]) if sa else []
    weights = await sig(conn, ORG_A, sa[0]["version_id"]) if sa else []
    check(len(sa) == 1 and not fu["errors"] and len(tk) == 1 and tk[0]["kind"] == "model_version_review"
          and [w[1] for w in weights] == [Decimal("0.55"), Decimal("0.45")] and int(nb) == 1,
          "a new platform version drafts the subscriber IN ITS OWN ORG with a review task; org B's copy is untouched",
          f"{fu}")


async def y_rep1(conn):
    section("[Y] rep1 approvals on human tasks, behaviour unchanged; verify_rep2 CSV check corrected")
    from services import replica_terms as T

    entry, ver = uuid.uuid4(), uuid.uuid4()
    async with conn.transaction():
        await conn.execute("SELECT set_config('app.is_super_admin', 'true', true)")
        src = await conn.fetchval("SELECT key FROM portfolio.replica_holdings_sources ORDER BY key LIMIT 1")
        cls = await conn.fetchval("SELECT key FROM portfolio.replica_instrument_classes WHERE enabled ORDER BY key LIMIT 1")
        await conn.execute("INSERT INTO portfolio.replica_securities (id, org_id, global_security_id, instrument_class, "
                           "ticker, holdings_source_key, display_name, created_by) VALUES ($1, $2, $3::uuid, $4, $5, $6, "
                           "$7, $8)", entry, ORG_A, S["sec_F"], cls, TICKER["F"], src, f"Holdings of {TICKER['F']}",
                           U_ADMIN1)
        await conn.execute("INSERT INTO portfolio.replica_terms_versions (id, org_id, replica_security_id, version_no, "
                           "maker_ids, created_by) VALUES ($1, $2, $3, 1, $4::uuid[], $5)", ver, ORG_A, entry, [U_ADMIN1],
                           U_ADMIN1)
    from services.database import get_pool

    pool = await get_pool()
    with _ctx(ORG_A):
        async with pool.acquire() as c:
            e = await T.load_entry(c, ORG_A, entry)
            v = dict(await c.fetchrow("SELECT * FROM portfolio.replica_terms_versions WHERE id = $1", ver))
        task = await T.ensure_approval_task(pool, org_id=ORG_A, entry=e, version=v, actor_id=U_ADMIN1)
    todos = await scoped(conn, ORG_A, "SELECT user_id, status FROM member_todos WHERE source = 'workflow_user_task' AND "
                                      "related_id = $1::uuid", task["id"]) if task else []
    old = await scoped(conn, ORG_A, "SELECT count(*) FROM member_todos WHERE source = 'replica_terms_approval'",
                       fetch="fetchval")
    check(task is not None and task["kind"] == "replica_terms_approval"
          and task["eligible_permission"] == "approve_replication"
          and {str(t["user_id"]) for t in todos} >= {str(U_ADMIN1), str(U_ADMIN2)} and int(old) == 0,
          "a terms draft opens a human task (approve_replication), notified through the engine; no old-path todos")
    path = f"/api/v1/replication/registry/{entry}/terms/{ver}/approve"
    c1, o1 = await api(U_ADMIN1, "POST", path, {"note": "ok"})
    check(c1 == 403 and o1.get("detail") == T.MAKER_CHECKER_RULE,
          "rep1 maker-checker refusal unchanged (403, the same text)", f"{c1} {o1}")
    c2, o2 = await api(U_ADMIN2, "POST", path, {"note": "ok"})
    t2 = await scoped(conn, ORG_A, "SELECT status, approved_by, result->>'outcome' AS o FROM workflow_run_steps WHERE "
                                   "id = $1::uuid", task["id"], fetch="fetchrow") if task else None
    acts = await scoped(conn, ORG_A, "SELECT count(*) FROM portfolio.replica_terms_actions WHERE terms_version_id = "
                                     "$1::uuid AND action = 'approve'", ver, fetch="fetchval")
    check(c2 == 200 and o2["version"]["status"] == "approved" and t2 is not None and t2["status"] == "completed"
          and t2["approved_by"] == U_ADMIN2 and t2["o"] == "approve" and int(acts) == 1,
          "the checker's approval completes the task; replica_terms_actions still logs it", f"{c2}")
    c3, _ = await api(U_ADMIN2, "POST", path, {"note": "again"})
    check(c3 == 409, "approving an approved version is still 409", f"{c3}")
    r1 = (API_DIR / "scripts" / "verify_rep1.py").read_text(encoding="utf-8")
    r2 = (API_DIR / "scripts" / "verify_rep2.py").read_text(encoding="utf-8")
    check("source = 'workflow_user_task'" in r1 and "approval human task is completed" in r1
          and "replica_terms_approval' AND related_id = $1\", UUID" not in r1,
          "verify_rep1.py asserts the human-task path (only its approval-task checks changed)")
    check("plan.exceptions" in r2 and 'c["unmatched"] == 1' not in r2,
          "verify_rep2.py reads the exception count from plan.exceptions (the importer is unchanged)")
    imp_now = (API_DIR / "services" / "custody" / "importer.py").read_text(encoding="utf-8")
    check(imp_now == git_show(BEFORE_REF, "apps/api/services/custody/importer.py"), "the importer is byte-identical")


async def y_isolation(conn):
    section("[Y] Org isolation both directions under app_service (rolbypassrls = false)")
    bypass = await conn.fetchval("SELECT rolbypassrls FROM pg_roles WHERE rolname = current_user")
    check(bypass is False, "isolation is tested on a role WITHOUT rolbypassrls", f"{await conn.fetchval('SELECT current_user')}")
    code, out, _ = await to_draft(U_ADMIN_B1, "b.csv", csv_bytes(["Ticker", "Weight"], [[TICKER["A"], 100]]),
                                  "REP3 Org B Model")
    mb = out.get("model", {}).get("id") if code == 201 else None
    ma = S.get("m_csv")
    check(bool(mb) and bool(ma), "each org has its own model", f"{code}")
    if not (mb and ma):
        return
    for viewer, own, other, label in ((U_ADMIN_B1, mb, ma, "B → A"), (U_ADMIN1, ma, mb, "A → B")):
        c_own, _ = await api(viewer, "GET", f"/api/v1/models/{own}")
        c_oth, _ = await api(viewer, "GET", f"/api/v1/models/{other}")
        c_wr, _ = await api(viewer, "POST", f"/api/v1/models/{other}/drafts", {})
        _, lst = await api(viewer, "GET", "/api/v1/models")
        _, log = await api(viewer, "GET", "/api/v1/models/uploads")
        ids = {r["id"] for r in lst.get("rows", [])}
        lorg = {r.get("model_id") for r in log.get("rows", [])}
        check(c_own == 200 and c_oth == 404 and c_wr == 404 and own in ids and other not in ids and other not in lorg,
              f"{label}: own 200; the other org's model 404 (read AND write), absent from list and upload log",
              f"{c_own} {c_oth} {c_wr}")
    for org, other in ((ORG_A, ORG_B), (ORG_B, ORG_A)):
        q = {t: int(await scoped(conn, org, f"SELECT count(*) FROM portfolio.{t} WHERE org_id = $1", other,
                                 fetch="fetchval"))
             for t in ("models", "model_versions", "model_version_positions", "model_uploads", "model_upload_events",
                       "model_events", "model_column_mappings", "model_subscriptions")}
        own = await scoped(conn, org, "SELECT count(*) FROM portfolio.models WHERE org_id = $1", org, fetch="fetchval")
        check(int(own) > 0 and all(v == 0 for v in q.values()),
              f"SQL as {str(org)[-4:]}: own models visible; every model table of the other org reads 0",
              f"own {own} other {q}")
    try:
        await scoped(conn, ORG_B, "INSERT INTO portfolio.model_events (org_id, action, actor_id) VALUES ($1, 'copy', $2)",
                     ORG_A, U_ADMIN_B1, fetch="execute")
        check(False, "org B cannot write an org A row")
    except Exception as exc:  # noqa: BLE001
        check("row-level security" in str(exc), "org B cannot write an org A row (RLS WITH CHECK)", str(exc)[:90])
    plat = await scoped(conn, ORG_B, "SELECT count(*) FROM portfolio.models WHERE id = $1::uuid", S.get("plat_m"),
                        fetch="fetchval") if S.get("plat_m") else None
    check(plat is not None and int(plat) == 1, "the platform entry's model is readable to another org (carve-out, read only)")


async def y_gates(conn):
    section("[Y] Gates both ways on identical requests")
    m = S.get("m_csv")
    c_v, o_v = await api(U_VIEWER, "GET", "/api/v1/models")
    c_a, o_a = await api(U_ADMIN1, "GET", "/api/v1/models")
    check(c_v == 200 and o_v["permissions"]["can_write"] is False and o_v["vocabularies"]["editable"] == []
          and o_v.get("actions") == [] and c_a == 200 and o_a["permissions"]["can_write"] is True
          and o_a["vocabularies"]["editable"] != [],
          "view_models reads with a view-only envelope (editable [], no actions); an admin's envelope allows writes")
    c_m, o_m = await api(U_MEMBER, "GET", "/api/v1/models")
    check(c_m == 403 and "view_models" in str(o_m.get("detail")), "no view_models → 403 naming it", f"{c_m}")
    cv, ov = await upload(U_VIEWER, "viewer.csv", csv_bytes(["Ticker", "Weight"], [[TICKER["B"], 100]]))
    ca, oa = await upload(U_ADMIN1, "viewer.csv", csv_bytes(["Ticker", "Weight"], [[TICKER["B"], 100]]))
    check(cv == 403 and "manage_models" in str(ov.get("detail")) and ca == 201,
          "the identical upload: viewer 403 (manage_models), admin 201", f"{cv} {ca}")
    if m:
        cd, _ = await api(U_ADMIN1, "POST", f"/api/v1/models/{m}/drafts", {})
        vid = (await scoped(conn, ORG_A, "SELECT id FROM portfolio.model_versions WHERE model_id = $1::uuid AND status = "
                                         "'draft' AND system_to IS NULL", m, fetch="fetchval"))
        if vid:
            await api(U_ADMIN1, "POST", f"/api/v1/models/{m}/versions/{vid}/submit", {})
            c1, o1 = await api(U_VIEWER, "POST", f"/api/v1/models/{m}/versions/{vid}/approve", {})
            c2, _ = await api(U_ADMIN2, "POST", f"/api/v1/models/{m}/versions/{vid}/approve", {})
            check(c1 == 403 and "approve_models" in str(o1.get("detail")) and c2 == 200,
                  "the identical approve: viewer 403 (approve_models), checker 200", f"{c1} {c2}")
        else:
            check(False, "a draft for the approve gate proof", f"{cd}")
    c3, _ = await api(U_VIEWER, "POST", "/api/v1/models/library", {"model_id": str(m), "shelf": "org", "title": "x",
                                                                    "provider": "x", "update_cadence": "monthly"})
    check(c3 == 403, "a viewer cannot publish (403)", f"{c3}")


async def y_new_org(conn):
    section("[Y] New org's first org_admin gets the three model permissions")
    from services import rbac

    before = await platform(conn, "SELECT count(*) FROM roles WHERE org_id = $1 AND name = 'org_admin'", ORG_C,
                            fetch="fetchval")
    async with conn.transaction():
        await conn.execute("SELECT set_config('app.is_super_admin', 'true', true)")
        await rbac.grant_org_admin(conn, U_NEWADMIN, ORG_C)
    dsn, _ = await admin_dsn()
    fresh = await connect(dsn)
    try:
        perms = {r["name"] for r in await platform(fresh, """
            SELECT p.name FROM roles r JOIN role_permissions rp ON rp.role_id = r.id JOIN permissions p
            ON p.id = rp.permission_id WHERE r.org_id = $1 AND r.name = 'org_admin'""", ORG_C)}
    finally:
        await fresh.close()
    check(int(before) == 0 and set(MODEL_PERMS) <= perms and "manage_org_settings" in perms,
          "a NEW org's first org_admin role carries view/manage/approve_models (re-read on an independent connection)",
          f"{sorted(perms)}")


async def y_routes(conn):
    section("[Y] Every new route gated and in AUTHZ1_ROUTES.json; --check CURRENT; no existing gate changed")
    before = json.loads(git_show(BEFORE_REF, "docs/AUTHZ1_ROUTES.json") or "{}")
    now = json.loads((REPO / "docs" / "AUTHZ1_ROUTES.json").read_text(encoding="utf-8"))

    def gates(d):
        return {(tuple(r.get("methods") or []), r["path"]): r for r in d.get("routes", [])}

    gb, gn = gates(before), gates(now)
    changed = [k for k in gb if k in gn and gb[k].get("after") != gn[k].get("after")]
    removed = [k for k in gb if k not in gn]
    check(len(gb) > 0 and not changed and not removed, "no existing route's gate changed or vanished",
          f"{changed[:5]} {removed[:5]}")
    # by path: later sprints add routes of their own (rep4: /replication/...) — rep3's are the /models ones
    new = [k for k in gn if k not in gb and k[1].startswith("/api/v1/models")]
    check(len(new) == 28 and all(any(p in str(gn[k].get("after")) for p in MODEL_PERMS) for k in new),
          "28 new routes under /models, each naming a model permission", f"{len(new)}")
    code, out, err = run_cmd([sys.executable, str(HERE / "authz1_routes.py"), "--check"], REPO, timeout=600)
    check(code == 0 and "CURRENT" in out, "authz1_routes.py --check: CURRENT", (out + err)[-200:])


async def y_front_end(conn):
    section("[Y] node --test tests/ passes; npm run build exits 0 (run by this script, not during the sprint)")
    code, out, err = run_cmd(["node", "--test", "tests/"], WEB_DIR, timeout=900)
    check(code == 0 and re.search(r"^# fail 0", out, re.M) is not None, "node --test tests/ passes (0 fail)",
          (out + err)[-300:] if code else "")
    check((WEB_DIR / "tests" / "models" / "rep3.test.mjs").exists() and "models list: no Upload" in out
          and "model page: a view-only envelope renders no write, approve or publish control" in out,
          "the rep3 node tests (no Upload / Approve / Publish / Edit controls without the flags) ran in that suite")
    if os.environ.get("SKIP_BUILD") == "1":
        print("[SKIP] npm run build (SKIP_BUILD=1 — a nested regression run; the outer verify builds)")
        return
    code, out, err = run_cmd(["npm", "run", "build"], WEB_DIR, timeout=1200)
    check(code == 0, "npm run build exits 0", (out + err)[-400:] if code else "")


def regression():
    section("[Y] Regression — verify_rep1.py ONCE, flat, AFTER teardown, VERIFY_NESTED=1")
    env = os.environ | {"VERIFY_NESTED": "1", "SKIP_BUILD": "1"}
    code, out, err = run_cmd([sys.executable, str(HERE / "verify_rep1.py")], REPO, timeout=1200, env=env)
    total = re.findall(r"^TOTAL: (\d+) PASS, (\d+) FAIL", out, re.M)
    result = re.findall(r"^RESULT: (.*)$", out, re.M)
    fails = [ln for ln in out.splitlines() if ln.lstrip().startswith("[FAIL]")]
    zero = (bool(total) and total[-1][1] == "0") or (bool(result) and "ALL ASSERTIONS PASSED" in result[-1])
    check(zero and not fails and code == 0, "verify_rep1.py: 0 FAIL ('TOTAL:' or 'RESULT:' summary)",
          f"{total[-1] if total else (result[-1] if result else 'no summary')}; {fails[:3]} "
          f"{err[-200:] if not (total or result) else ''}")


async def run(name, fn, *args):
    try:
        res = fn(*args)
        if asyncio.iscoroutine(res):
            await res
    except Exception as exc:  # noqa: BLE001
        check(False, f"{name} raised", f"{type(exc).__name__}: {exc}")
        traceback.print_exc()


async def guard_prefix(conn) -> bool:
    n = await platform(conn, "SELECT count(*) FROM portfolio.securities_global_identifiers i JOIN "
                             "portfolio.securities_global s ON s.id = i.global_security_id WHERE i.id_value LIKE $1 "
                             "AND NOT (s.name = ANY($2::text[]))", f"{PREFIX}%", list(SEC_NAMES), fetch="fetchval")
    m = await platform(conn, "SELECT count(*) FROM portfolio.models WHERE name LIKE 'REP3 %' AND NOT (org_id = "
                             "ANY($1::uuid[])) AND (created_by IS NULL OR created_by <> $2)",
                       [str(o) for o in FIXTURE_ORGS], U_PLAT, fetch="fetchval")
    return check(int(n) == 0 and int(m) == 0, f"the reserved prefixes ({PREFIX}, 'REP3 ' model names) match no real row",
                 f"{n}/{m}")


DROPPED_CONNECTION = (asyncpg.ConnectionDoesNotExistError, asyncpg.InterfaceError, ConnectionError)


async def main_async() -> int:
    dsn, how = await admin_dsn()
    if not dsn:
        print(f"[FAIL] no working DATABASE_URL: {how}")
        print("TOTAL: 0 PASS, 1 FAIL")
        return 1
    os.environ["DATABASE_URL"] = dsn
    print(f"[INFO] database via {how}")
    conn = await connect(dsn)
    install_stubs()
    pre = real = None
    try:
        if not await guard_prefix(conn):
            print(f"\nTOTAL: {_n_pass} PASS, {_n_fail} FAIL")
            return 1
        await run("pre-teardown", teardown, conn)
        pre = await counts(conn)
        real = await real_fingerprints(conn)
        await run("task 1", y_task1, conn)
        await run("seed", seed, conn)
        if S.get("sec_A"):
            await run("schema", y_schema, conn)
            await run("upload", y_upload, conn)
            await run("mapping", y_mapping, conn)
            await run("log", y_log, conn)
            await run("versions", y_versions, conn)
            await run("names", y_names, conn)
            await run("sleeves", y_sleeves, conn)
            await run("library", y_library, conn)
            await run("rep1", y_rep1, conn)
            await run("isolation", y_isolation, conn)
            await run("gates", y_gates, conn)
            await run("new org", y_new_org, conn)
        else:
            check(False, "seed created the fixture securities", "every fixture-scoped section skipped")
        await run("routes", y_routes, conn)
        await run("front end", y_front_end, conn)
        check(NETWORK_CALLS == [], "no network call was attempted during the whole run", f"{NETWORK_CALLS[:3]}")
    finally:
        section("[Y] Teardown — zero fixture rows; every touched table at its pre-test count; real rows unchanged")
        try:
            from services.database import close_pool

            await close_pool()
        except Exception:  # noqa: BLE001
            pass

        async def after_teardown(c):
            """Every post-teardown read, as ONE unit, so a retry re-runs all of it on the same connection."""
            return (await leftovers(c), await counts(c) if pre is not None else None,
                    await real_fingerprints(c) if real is not None else None)

        try:
            log = await teardown(conn)
            print(f"[INFO] teardown removed: {log}")
            try:
                left, post, after = await after_teardown(conn)
            except DROPPED_CONNECTION as exc:
                # Reconnect ONCE and re-run every count on the new connection. A second drop propagates
                # to "teardown raised" below — a dropped connection is never reported as clean.
                print(f"[INFO] connection dropped during the leftover check ({type(exc).__name__}); reconnecting once")
                try:
                    conn.terminate()
                except Exception:  # noqa: BLE001
                    pass
                conn = await connect(dsn)
                left, post, after = await after_teardown(conn)
            check(left == 0, "zero fixture rows remain", f"{left}")
            if post is not None:
                for t in COUNTED:
                    check(post[t] == pre[t], f"{t} at its pre-test count", f"{pre[t]} → {post[t]}")
            if after is not None:
                for k in real:
                    check(after[k] == real[k], f"real {k} unchanged (count + checksum)", f"{real[k][0]} → {after[k][0]}")
        except Exception as exc:  # noqa: BLE001
            check(False, "teardown raised", f"{type(exc).__name__}: {exc}")
            traceback.print_exc()
        try:
            await conn.close()
        except Exception:  # noqa: BLE001
            pass
    if os.environ.get("VERIFY_NESTED") == "1":
        print("[SKIP] nested run")
    else:
        try:
            regression()
        except Exception as exc:  # noqa: BLE001
            check(False, "regression raised", f"{type(exc).__name__}: {exc}")
    print(f"\nTOTAL: {_n_pass} PASS, {_n_fail} FAIL")
    return 0 if _n_fail == 0 else 1


if __name__ == "__main__":
    sys.exit(asyncio.run(main_async()))
